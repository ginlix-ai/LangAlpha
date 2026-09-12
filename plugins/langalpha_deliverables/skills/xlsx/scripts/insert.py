#!/usr/bin/env python3
"""Insert or delete worksheet rows or columns and repair workbook references.

Usage:
    python insert.py rows <file.xlsx> --sheet NAME --at N [--count K] [--copy-style-from ROW] [--out PATH] [--dry-run]
    python insert.py columns <file.xlsx> --sheet NAME --at LETTER [--count K] [--copy-style-from LETTER] [--out PATH] [--dry-run]
    python insert.py delete-rows <file.xlsx> --sheet NAME --at N [--count K] [--out PATH] [--dry-run]
    python insert.py delete-columns <file.xlsx> --sheet NAME --at LETTER [--count K] [--out PATH] [--dry-run]

--at is a positive, 1-based row number or a column letter (A through XFD).
--count defaults to 1. Insertions occur before --at; deletions start at --at.
--copy-style-from uses the source's original position and copies cell style
indices and row height or column width into blank inserted cells. Table
columns also receive unique placeholder headers, reported in warnings.
--out writes a separate file; otherwise the input is replaced atomically.
--dry-run computes and checks the complete result in memory, writes nothing,
and adds dry_run: true to the report. --help prints this full usage.

References are changed in XML package parts, without an openpyxl round trip.
Ranges grow only when an insertion is inside them, shrink on partial deletion,
and become #REF! when deleted entirely. External workbook links and pivot
cache sources remain unchanged and are reported. Shared formulas are retained
when their translations remain consistent; otherwise they become explicit
formulas. Calculation chains are removed and full calculation is requested.
Run recalc.py afterwards to rebuild cached values.

Output is one JSON object on stdout:
    status              success or error (only error exits 1)
    sheet, operation    target and insert_rows, insert_columns,
                        delete_rows or delete_columns
    at, count           insertion/deletion position as supplied and count
    cells_shifted       surviving cells whose coordinates changed
    formulas_rewritten  changed effective cell or other formula texts
    names_rewritten     changed defined-name expressions
    charts_rewritten    chart parts with changed formula references
    parts_changed       package parts written or deliberately removed
    warnings            located #REF! occurrences, external links, dropped
                        features, placeholder headers and other review items
    self_check          {ok: true} or {ok: false, problems: [...]}

Existing #REF! errors are included in warnings so the self-check can reconcile
all error occurrences, excluding quoted text. Shared followers are checked as
effective formulas. Sheet-local names use their local sheet; unqualified
workbook names use the saved active sheet. Deleted drawing endpoints clamp to
the deletion boundary; selections use surviving areas, or that boundary when
none survive. A failed self-check leaves the input and any output untouched.
"""

from __future__ import annotations

import copy
import io
import json
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter

SML = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DRAW = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
VML = "urn:schemas-microsoft-com:vml"
EXCEL = "urn:schemas-microsoft-com:office:excel"
MAX_ROW, MAX_COL = 1048576, 16384
CELL = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)([1-9][0-9]*)\Z")
COL = re.compile(r"(\$?)([A-Za-z]{1,3})\Z")
ROW = re.compile(r"(\$?)([1-9][0-9]*)\Z")
FORMULA_TAGS = {
    "formula",
    "formula1",
    "formula2",
    "calculatedColumnFormula",
    "totalsRowFormula",
}


def tag(name):
    return f"{{{SML}}}{name}"


def local(node):
    return etree.QName(node).localname if isinstance(node.tag, str) else ""


def parse_xml(data):
    return etree.fromstring(
        data, etree.XMLParser(resolve_entities=False, no_network=True)
    )


def relation_path(part):
    return posixpath.join(
        posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels"
    )


def resolve_part(owner, target):
    return (
        target.lstrip("/")
        if target.startswith("/")
        else posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
    )


def endpoint(value, kind):
    match = {"cell": CELL, "column": COL, "row": ROW}[kind].fullmatch(value)
    if not match:
        raise ValueError(f"Invalid {kind} reference: {value!r}")
    p = match.groups()
    col = column_index_from_string(p[1]) if kind != "row" else None
    row = int(p[3] if kind == "cell" else p[1]) if kind != "column" else None
    if (col is not None and col > MAX_COL) or (row is not None and row > MAX_ROW):
        raise ValueError(f"Reference outside Excel grid: {value!r}")
    return p, col, row


def parse_range(value):
    bits = value.split(":")
    if len(bits) == 1:
        kind = "cell"
    elif len(bits) == 2:
        kind = (
            "cell"
            if CELL.fullmatch(bits[0])
            else "column"
            if COL.fullmatch(bits[0])
            else "row"
        )
    else:
        raise ValueError(f"Invalid range: {value!r}")
    return kind, [endpoint(bit, kind) for bit in bits]


def xy(value):
    _, col, row = endpoint(value, "cell")
    return col, row


def address(col, row):
    return f"{get_column_letter(col)}{row}"


def formula_tokens(text):
    return Tokenizer(text if text.startswith("=") else "=" + text).items


def error_count(text):
    return sum(
        t.value.count("#REF!")
        for t in formula_tokens(text)
        if t.type == "OPERAND" and t.subtype in {"RANGE", "ERROR"}
    )


def validate_formula(text, defined_name=False):
    stack = []
    for token in formula_tokens(text):
        if token.subtype == "OPEN" and token.type in {"FUNC", "PAREN", "ARRAY"}:
            stack.append(token.type)
        elif token.subtype == "CLOSE" and token.type in {"FUNC", "PAREN", "ARRAY"}:
            if not stack or stack.pop() != token.type:
                raise ValueError("Mismatched formula delimiters")
        if defined_name and token.type == "OPERAND" and token.subtype == "RANGE":
            value = token.value
            if "#REF!" in value or "[" in value:
                continue
            ref = value.rsplit("!", 1)[-1]
            if ":" in ref or "$" in ref:
                parse_range(ref)
    if stack:
        raise ValueError("Unclosed formula delimiters")


class Shift:
    def __init__(self, sheet, sheets, axis, at, count, delete):
        self.sheet = sheet
        self.sheets = [s.casefold() for s in sheets]
        self.axis, self.at, self.count, self.delete = axis, at, count, delete
        self.limit = MAX_ROW if axis == "row" else MAX_COL

    def point(self, n, clamp=False):
        if n < self.at:
            return n
        if not self.delete:
            return n + self.count
        if n < self.at + self.count:
            return min(self.at, self.limit) if clamp else None
        return n - self.count

    def interval(self, lo, hi):
        reverse = lo > hi
        if reverse:
            lo, hi = hi, lo
        if not self.delete:
            if lo >= self.at:
                lo += self.count
            if hi >= self.at:
                hi += self.count
            if lo > self.limit:
                return None
            hi = min(hi, self.limit)
        else:
            end = self.at + self.count - 1
            if lo >= self.at and hi <= end:
                return None
            lo = lo if lo < self.at else self.at if lo <= end else lo - self.count
            hi = hi if hi < self.at else self.at - 1 if hi <= end else hi - self.count
        return (hi, lo) if reverse else (lo, hi)

    def range(self, value):
        kind, ends = parse_range(value)
        if (kind == "column" and self.axis == "row") or (
            kind == "row" and self.axis == "column"
        ):
            return value
        index = 2 if self.axis == "row" else 1
        nums = [e[index] for e in ends]
        moved = self.interval(nums[0], nums[-1])
        if moved is None:
            return None
        if nums == list(moved) or (len(nums) == 1 and nums[0] == moved[0]):
            return value
        result = []
        for (p, col, row), n in zip(ends, moved):
            if kind == "cell":
                col, row = (col, n) if self.axis == "row" else (n, row)
                letters = get_column_letter(col)
                if p[1].islower():
                    letters = letters.lower()
                result.append(f"{p[0]}{letters}{p[2]}{row}")
            elif kind == "column":
                letters = get_column_letter(n)
                result.append(p[0] + (letters.lower() if p[1].islower() else letters))
            else:
                result.append(p[0] + str(n))
        return ":".join(result)

    def coordinate(self, ref, clamp=False):
        col, row = xy(ref)
        n = self.point(row if self.axis == "row" else col, clamp)
        if n is None:
            return None
        if n > self.limit:
            raise ValueError("The edit would move cells outside Excel's grid")
        return address(col, n) if self.axis == "row" else address(n, row)

    def in_scope(self, qualifier, context):
        if qualifier is None:
            return context is not None and context.casefold() == self.sheet.casefold()
        if "[" in qualifier:
            return False
        # Split outside quotes first, then allow Excel's single quoted 3D span.
        names, start, quoted, i = [], 0, False, 0
        while i < len(qualifier):
            char = qualifier[i]
            if char == "'":
                if quoted and i + 1 < len(qualifier) and qualifier[i + 1] == "'":
                    i += 2
                    continue
                quoted = not quoted
            elif char == ":" and not quoted:
                names.append(qualifier[start:i])
                start = i + 1
            i += 1
        names.append(qualifier[start:])
        names = [
            (s[1:-1].replace("''", "'") if s.startswith("'") and s.endswith("'") else s)
            for s in names
        ]
        if len(names) == 1:
            names = names[0].split(":")
        names = [s.casefold() for s in names]
        if len(names) == 1:
            return names[0] == self.sheet.casefold()
        if len(names) != 2 or any(s not in self.sheets for s in names):
            return False
        a, b = sorted(self.sheets.index(s) for s in names)
        return a <= self.sheets.index(self.sheet.casefold()) <= b

    def operand(self, value, context):
        if "#REF!" in value:
            return value
        qualifier, ref = value.rsplit("!", 1) if "!" in value else (None, value)
        if not self.in_scope(qualifier, context):
            return value
        try:
            parse_range(ref)
        except ValueError:
            return value  # RANGE tokens also contain names and structured references.
        moved = self.range(ref)
        return (qualifier + "!" if qualifier is not None else "") + (moved or "#REF!")

    def formula(self, text, context):
        if not text:
            return text
        tokens = formula_tokens(text)
        cursor = 1 if text.startswith("=") else 0
        edits = []
        for token in tokens:
            if token.type == "WHITE-SPACE":
                while cursor < len(text) and text[cursor].isspace():
                    cursor += 1
                continue
            start = text.find(token.value, cursor)
            if start < 0:
                raise ValueError(f"Cannot locate formula token in {text!r}")
            cursor = start + len(token.value)
            if token.type == "OPERAND" and token.subtype == "RANGE":
                changed = self.operand(token.value, context)
                if changed != token.value:
                    edits.append((start, cursor, changed))
        for start, end, value in reversed(edits):
            text = text[:start] + value + text[end:]
        return text


class Package:
    def __init__(self, source):
        with zipfile.ZipFile(source) as archive:
            bad = archive.testzip()
            if bad:
                raise ValueError(f"Corrupt zip member: {bad}")
            self.infos = archive.infolist()
            self.comment = archive.comment
            self.original = {i.filename: archive.read(i.filename) for i in self.infos}
        if len(self.infos) != len(self.original):
            raise ValueError("Duplicate zip member names")
        for required in (
            "[Content_Types].xml",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
        ):
            if required not in self.original:
                raise ValueError(f"Not a workbook package: missing {required}")
        self.trees, self.before, self.removed = {}, {}, set()
        self.workbook = self.xml("xl/workbook.xml")
        rels = {
            r.get("Id"): resolve_part("xl/workbook.xml", r.get("Target", ""))
            for r in self.xml("xl/_rels/workbook.xml.rels")
            if r.get("TargetMode") != "External"
        }
        self.sheets = {
            s.get("name"): rels[s.get(f"{{{REL}}}id")]
            for s in self.workbook.iter(tag("sheet"))
        }
        if not self.sheets:
            raise ValueError("Workbook has no sheets")

    def xml(self, part):
        if part not in self.trees:
            self.trees[part] = parse_xml(self.original[part])
            self.before[part] = etree.tostring(self.trees[part])
        return self.trees[part]

    def relationships(self, owner):
        part = relation_path(owner)
        if part not in self.original or part in self.removed:
            return []
        return [
            (r, resolve_part(owner, r.get("Target", "")))
            for r in self.xml(part)
            if r.get("TargetMode") != "External"
        ]

    def owners(self):
        owners = {}
        for sheet, part in self.sheets.items():
            pending, seen = [part], set()
            while pending:
                part = pending.pop()
                if part in seen:
                    continue
                seen.add(part)
                owners.setdefault(part, sheet)
                pending.extend(
                    p for _, p in self.relationships(part) if p in self.original
                )
        return owners

    def discard(self, part):
        if part in self.original:
            self.removed.add(part)
        relpart = relation_path(part)
        if relpart in self.original:
            self.removed.add(relpart)
        types = self.xml("[Content_Types].xml")
        for child in list(types):
            if child.get("PartName", "").lstrip("/") in self.removed:
                types.remove(child)

    def payload(self):
        payload = {p: b for p, b in self.original.items() if p not in self.removed}
        for part, root in self.trees.items():
            if part not in self.removed and etree.tostring(root) != self.before[part]:
                payload[part] = etree.tostring(
                    root, encoding="UTF-8", xml_declaration=True
                )
        return payload

    def write(self, destination, payload):
        with zipfile.ZipFile(destination, "w") as archive:
            archive.comment = self.comment
            for info in self.infos:
                if info.filename in payload:
                    archive.writestr(info, payload[info.filename])


def cell_formulas(root):
    cells = [(c, c.find(tag("f"))) for c in root.iter(tag("c"))]
    masters = {
        f.get("si"): (c.get("r"), f.text)
        for c, f in cells
        if f is not None and f.get("t") == "shared" and f.text
    }
    result = []
    for cell, formula in cells:
        if formula is None:
            continue
        text = formula.text or ""
        if formula.get("t") == "shared" and not text:
            si = formula.get("si")
            if si not in masters:
                raise ValueError(
                    f"Missing shared formula master {si} at {cell.get('r')}"
                )
            origin, master = masters[si]
            text = Translator("=" + master, origin=origin).translate_formula(
                cell.get("r")
            )[1:]
        result.append((cell, formula, text))
    return result


def ensure_cell(root, ref):
    col, row = xy(ref)
    data = root.find(tag("sheetData"))
    if data is None:
        raise ValueError("Target sheet has no sheetData")
    row_node = next((r for r in data if int(r.get("r")) == row), None)
    if row_node is None:
        row_node = etree.Element(tag("row"), r=str(row))
        before = next((r for r in data if int(r.get("r")) > row), None)
        if before is not None:
            before.addprevious(row_node)
        else:
            data.append(row_node)
    cell = next((c for c in row_node if c.get("r") == ref), None)
    if cell is None:
        cell = etree.Element(tag("c"), r=ref)
        before = next(
            (c for c in row_node if local(c) == "c" and xy(c.get("r"))[0] > col), None
        )
        if before is not None:
            before.addprevious(cell)
        else:
            row_node.append(cell)
    return cell


def set_header(root, ref, name):
    cell = ensure_cell(root, ref)
    for child in list(cell):
        if local(child) in {"v", "f", "is"}:
            cell.remove(child)
    cell.set("t", "inlineStr")
    inline = etree.SubElement(cell, tag("is"))
    etree.SubElement(inline, tag("t")).text = name


class Editor:
    def __init__(self, package, sheet, axis, at, count, delete, copy_from, report):
        self.p, self.report = package, report
        self.shift = Shift(sheet, package.sheets, axis, at, count, delete)
        self.sheet, self.part = sheet, package.sheets[sheet]
        self.root = package.xml(self.part)
        if local(self.root) != "worksheet":
            raise ValueError("The target must be a worksheet")
        self.copy_from = copy_from
        self.owners = package.owners()

    def warn(self, kind, **details):
        self.report["warnings"].append({"type": kind, "sheet": self.sheet, **details})

    def text_formula(self, node, context, counter="formulas_rewritten"):
        old = node.text
        if old:
            node.text = self.shift.formula(old, context)
            if node.text != old:
                self.report[counter] += 1
                return True
        return False

    def rewrite_cells(self, root, sheet):
        records = cell_formulas(root)
        groups, arrays = {}, []
        for cell, formula, old in records:
            old_ref = cell.get("r")
            new_ref = self.shift.coordinate(old_ref) if sheet == self.sheet else old_ref
            new = self.shift.formula(old, sheet)
            if new_ref is not None and new != old:
                self.report["formulas_rewritten"] += 1
            if formula.get("t") == "shared":
                groups.setdefault(formula.get("si"), []).append(
                    (cell, formula, new_ref, new)
                )
            elif (
                formula.get("t") == "array"
                and formula.get("ref")
                and sheet == self.sheet
            ):
                if (
                    new_ref is None
                    and self.shift.range(formula.get("ref")) is not None
                    and new != old
                ):
                    self.report["formulas_rewritten"] += 1
                arrays.append(
                    (formula, self.shift.range(formula.get("ref")), new_ref, new)
                )
            else:
                if old:
                    formula.text = new
                if sheet == self.sheet:
                    for attr in ("ref", "r1", "r2"):
                        if formula.get(attr):
                            formula.set(
                                attr, self.shift.range(formula.get(attr)) or "#REF!"
                            )
        return groups, arrays

    def finish_formulas(self, groups, arrays, root, target, sheet):
        for si, entries in groups.items():
            survivors = [e for e in entries if e[2] is not None]
            if not survivors:
                continue
            survivors.sort(key=lambda e: xy(e[2])[::-1])
            master = survivors[0]
            old_master = next((e for e in entries if e[1].text), None)
            old_range = old_master[1].get("ref") if old_master else None
            new_range = (
                self.shift.range(old_range) if target and old_range else old_range
            )
            consistent = True
            for _, _, ref, text in survivors:
                try:
                    translated = Translator(
                        "=" + master[3], origin=master[2]
                    ).translate_formula(ref)[1:]
                    consistent &= translated == text
                except Exception:
                    consistent = False
            if not consistent:
                self.warn(
                    "shared_formula_expanded",
                    sheet=sheet,
                    cell=master[2],
                    shared_index=si,
                    message="Structural edit requires independent shared follower formulas",
                )
            for entry in survivors:
                _, formula, ref, text = entry
                if consistent:
                    formula.text = text if entry is master else None
                    if entry is master:
                        formula.set("ref", new_range or ref)
                    else:
                        formula.attrib.pop("ref", None)
                else:
                    for attr in ("t", "si", "ref"):
                        formula.attrib.pop(attr, None)
                    formula.text = text
        for formula, new_range, new_ref, text in arrays:
            if new_range is None:
                self.warn(
                    "array_formula_dropped",
                    message="The entire array formula range was deleted",
                )
                continue
            formula.set("ref", new_range)
            formula.text = text
            if new_ref is None:
                new_ref = new_range.split(":")[0].replace("$", "")
                cell = ensure_cell(root, new_ref)
                for child in list(cell):
                    if local(child) in {"f", "v", "is"}:
                        cell.remove(child)
                cell.attrib.pop("t", None)
                cell.insert(0, copy.deepcopy(formula))
                self.warn(
                    "array_master_promoted",
                    cell=new_ref,
                    message="The array master was deleted; retained the formula at the surviving range origin",
                )

    def shift_grid(self):
        shift, root = self.shift, self.root
        data = root.find(tag("sheetData"))
        if data is None:
            raise ValueError("Target sheet has no sheetData")
        used = [
            xy(c.get("r"))[1 if shift.axis == "row" else 0] for c in data.iter(tag("c"))
        ]
        if shift.axis == "row":
            used.extend(int(r.get("r")) for r in data)
        cols = root.find(tag("cols"))
        if shift.axis == "column" and cols is not None:
            used.extend(int(c.get("max")) for c in cols)
        if shift.delete and shift.at + shift.count - 1 > max(used, default=1):
            raise ValueError("Deletion reaches past the used range")
        source_cells, source_meta = [], {}
        if self.copy_from is not None:
            for row in data:
                if shift.axis == "row" and int(row.get("r")) == self.copy_from:
                    source_meta = {
                        k: row.get(k)
                        for k in ("ht", "customHeight")
                        if row.get(k) is not None
                    }
                for cell in row:
                    if local(cell) != "c":
                        continue
                    col_n, row_n = xy(cell.get("r"))
                    if (row_n if shift.axis == "row" else col_n) == self.copy_from:
                        source_cells.append((col_n, row_n, cell.get("s")))
            if shift.axis == "column" and cols is not None:
                for col in cols:
                    if int(col.get("min")) <= self.copy_from <= int(col.get("max")):
                        source_meta = {
                            k: col.get(k)
                            for k in ("width", "customWidth")
                            if col.get(k) is not None
                        }
        for row in list(data):
            row_n = int(row.get("r"))
            moved = shift.point(row_n) if shift.axis == "row" else row_n
            if moved is None:
                data.remove(row)
                continue
            if moved > MAX_ROW:
                raise ValueError("The edit would move rows outside Excel's grid")
            row.set("r", str(moved))
            for cell in list(row):
                if local(cell) != "c":
                    continue
                old = cell.get("r")
                new = shift.coordinate(old)
                if new is None:
                    row.remove(cell)
                elif old != new:
                    cell.set("r", new)
                    self.report["cells_shifted"] += 1
            if shift.axis == "column" and row.get("spans"):
                spans = []
                for span in row.get("spans").split():
                    a, b = map(int, span.split(":"))
                    moved_span = shift.interval(a, b)
                    if moved_span:
                        spans.append(f"{moved_span[0]}:{moved_span[1]}")
                if spans:
                    row.set("spans", " ".join(spans))
                else:
                    row.attrib.pop("spans", None)
        if shift.axis == "column" and cols is not None:
            for col in list(cols):
                lo, hi = int(col.get("min")), int(col.get("max"))
                if not shift.delete and lo < shift.at <= hi:
                    right = copy.deepcopy(col)
                    col.set("max", str(shift.at - 1))
                    right.set("min", str(shift.at + shift.count))
                    right.set("max", str(min(MAX_COL, hi + shift.count)))
                    if int(right.get("min")) <= MAX_COL:
                        col.addnext(right)
                else:
                    moved = shift.interval(lo, hi)
                    if moved is None:
                        cols.remove(col)
                    else:
                        col.set("min", str(moved[0]))
                        col.set("max", str(moved[1]))
        if self.copy_from is not None:
            for n in range(shift.at, shift.at + shift.count):
                if shift.axis == "row":
                    for col_n, _, style in source_cells:
                        cell = ensure_cell(root, address(col_n, n))
                        if style is not None:
                            cell.set("s", style)
                    row = next((r for r in data if int(r.get("r")) == n), None)
                    if row is None:
                        row = etree.Element(tag("row"), r=str(n))
                        data.append(row)
                    row.attrib.update(source_meta)
                else:
                    for _, row_n, style in source_cells:
                        cell = ensure_cell(root, address(n, row_n))
                        if style is not None:
                            cell.set("s", style)
                    if source_meta:
                        if cols is None:
                            cols = etree.Element(tag("cols"))
                            data.addprevious(cols)
                        etree.SubElement(
                            cols, tag("col"), min=str(n), max=str(n), **source_meta
                        )
        data[:] = sorted(data, key=lambda r: int(r.get("r")))
        if cols is not None:
            cols[:] = sorted(cols, key=lambda c: int(c.get("min")))
            if not len(cols):
                root.remove(cols)

    def range_attribute(self, node, attr, clamp=False):
        original = node.get(attr)
        if not original:
            return True
        values = [self.shift.range(v) for v in original.split()]
        moved = " ".join(v for v in values if v)
        if not moved and clamp:
            moved = self.shift.coordinate(
                original.split()[0].split(":")[0].replace("$", ""), True
            )
        if not moved:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
            self.warn(
                "dropped_range",
                element=local(node),
                ref=original,
                message="The entire referenced range was deleted",
            )
            return False
        node.set(attr, moved)
        return True

    def autofilter(self, node):
        old = node.get("ref")
        if not old:
            return
        old_left = xy(old.split(":")[0].replace("$", ""))[0]
        if not self.range_attribute(node, "ref"):
            return
        new_left = xy(node.get("ref").split(":")[0].replace("$", ""))[0]
        if self.shift.axis == "column":
            for child in list(node):
                if local(child) == "filterColumn":
                    moved = self.shift.point(old_left + int(child.get("colId")))
                    if moved is None:
                        node.remove(child)
                        self.warn(
                            "filter_column_dropped",
                            ref=old,
                            column_id=child.get("colId"),
                        )
                    else:
                        child.set("colId", str(moved - new_left))

    def sheet_ranges(self):
        root = self.root
        for node in list(root.iter()):
            kind = local(node)
            if kind in {"c", "row", "f", "dimension"} or not isinstance(node.tag, str):
                continue
            if kind == "autoFilter":
                self.autofilter(node)
                continue
            for attr in ("ref", "sqref"):
                if node.get(attr) and not self.range_attribute(
                    node, attr, kind == "selection"
                ):
                    break
            if kind == "selection" and node.get("activeCell"):
                active = self.shift.coordinate(node.get("activeCell"))
                if active is None:
                    active = (
                        node.get("sqref", "").split()[0].split(":")[0].replace("$", "")
                        if node.get("sqref")
                        else self.shift.coordinate(node.get("activeCell"), True)
                    )
                node.set("activeCell", active)
                # Deleting one area changes indices in a multi-area selection.
                if node.get("activeCellId") and int(node.get("activeCellId")) >= len(
                    node.get("sqref", "A1").split()
                ):
                    node.set("activeCellId", "0")
            if kind == "pane":
                old_x = float(node.get("xSplit", "0"))
                old_y = float(node.get("ySplit", "0"))
                if node.get("topLeftCell"):
                    node.set(
                        "topLeftCell",
                        self.shift.coordinate(node.get("topLeftCell"), True),
                    )
                attr = "ySplit" if self.shift.axis == "row" else "xSplit"
                if node.get("state") in {"frozen", "frozenSplit"} and node.get(attr):
                    split = int(float(node.get(attr)))
                    new_split = self.shift.point(split + 1, True) - 1
                    if new_split:
                        node.set(attr, str(new_split))
                    else:
                        node.attrib.pop(attr, None)
                if node.get("state") in {"frozen", "frozenSplit"}:
                    x, y = (
                        float(node.get("xSplit", "0")),
                        float(node.get("ySplit", "0")),
                    )
                    pane_name = (
                        "bottomRight"
                        if x and y
                        else "topRight"
                        if x
                        else "bottomLeft"
                        if y
                        else None
                    )
                    active = node.get("activePane")
                    if pane_name and (
                        active
                        not in {"topLeft", "topRight", "bottomLeft", "bottomRight"}
                        or (not x and active in {"topRight", "bottomRight"})
                        or (not y and active in {"bottomLeft", "bottomRight"})
                    ):
                        node.set("activePane", pane_name)
                    elif not pane_name:
                        node.getparent().remove(node)
                    for selection in root.iter(tag("selection")):
                        selected = selection.get("pane")
                        if old_x and not x:
                            selected = {
                                "topRight": None,
                                "bottomRight": "bottomLeft",
                            }.get(selected, selected)
                        if old_y and not y:
                            selected = {
                                "bottomLeft": None,
                                "bottomRight": "topRight",
                            }.get(selected, selected)
                        if selected:
                            selection.set("pane", selected)
                        else:
                            selection.attrib.pop("pane", None)
            if kind == "sqref" and node.text:
                moved = " ".join(
                    v for v in (self.shift.range(s) for s in node.text.split()) if v
                )
                if moved:
                    node.text = moved
                else:
                    parent = node.getparent()
                    if local(parent) == "sparkline":
                        parent.getparent().remove(parent)
                        self.warn("sparkline_dropped", ref=node.text)
                    else:
                        owner = parent
                        if owner.getparent() is not None:
                            owner.getparent().remove(owner)
                        self.warn("dropped_range", element=local(parent), ref=node.text)
        for node in list(root.iter()):
            kind = local(node)
            if kind in {"mergeCells", "dataValidations"}:
                if len(node):
                    node.set("count", str(len(node)))
                elif node.getparent() is not None:
                    node.getparent().remove(node)
            elif kind == "sparklineGroup" and not any(
                local(n) == "sparkline" for n in node.iter()
            ):
                node.getparent().remove(node)
        for view in root.iter(tag("sheetView")):
            seen = set()
            for selection in reversed(view.findall(tag("selection"))):
                pane = selection.get("pane", "topLeft")
                if pane in seen:
                    view.remove(selection)
                seen.add(pane)

    def tables(self):
        for relation, part in list(self.p.relationships(self.part)):
            if not relation.get("Type", "").endswith("/table"):
                continue
            table = self.p.xml(part)
            old = table.get("ref")
            new = self.shift.range(old)
            if new is None:
                self.p.discard(part)
                relation.getparent().remove(relation)
                for tp in list(self.root.iter(tag("tablePart"))):
                    if tp.get(f"{{{REL}}}id") == relation.get("Id"):
                        tp.getparent().remove(tp)
                self.warn("table_dropped", table=table.get("name"), ref=old)
                continue
            table.set("ref", new)
            left, top = xy(old.split(":")[0])
            right, bottom = xy(old.split(":")[-1])
            new_left, new_top = xy(new.split(":")[0])
            columns = table.find(tag("tableColumns"))
            if columns is None or len(columns) != right - left + 1:
                raise ValueError(f"Invalid table column list in {part}")
            if self.shift.axis == "column":
                if self.shift.delete:
                    for offset, column in enumerate(list(columns)):
                        if self.shift.point(left + offset) is None:
                            columns.remove(column)
                elif left < self.shift.at <= right:
                    names = {c.get("name", "").casefold() for c in columns}
                    next_id = max(int(c.get("id")) for c in columns) + 1
                    for offset in range(self.shift.count):
                        number = 1
                        while f"column{number}" in names:
                            number += 1
                        name = f"Column{number}"
                        names.add(name.casefold())
                        column = etree.Element(
                            tag("tableColumn"), id=str(next_id + offset), name=name
                        )
                        columns.insert(self.shift.at - left + offset, column)
                        if int(table.get("headerRowCount", "1")):
                            set_header(
                                self.root,
                                address(self.shift.at + offset, new_top),
                                name,
                            )
                        self.warn(
                            "table_header",
                            table=table.get("name"),
                            cell=address(self.shift.at + offset, new_top),
                            header=name,
                            message="Replace this placeholder with a meaningful table header",
                        )
                columns.set("count", str(len(columns)))
            elif self.shift.delete:
                if self.shift.point(top) is None and int(
                    table.get("headerRowCount", "1")
                ):
                    names = set()
                    strings = (
                        self.p.xml("xl/sharedStrings.xml")
                        if "xl/sharedStrings.xml" in self.p.original
                        else None
                    )
                    for offset, column in enumerate(columns):
                        ref = address(new_left + offset, new_top)
                        cell = ensure_cell(self.root, ref)
                        value = cell.findtext(tag("v"))
                        if (
                            cell.get("t") == "s"
                            and strings is not None
                            and value is not None
                        ):
                            value = "".join(strings[int(value)].itertext())
                        elif cell.get("t") == "inlineStr":
                            value = "".join(cell.find(tag("is")).itertext())
                        name = value or column.get("name") or f"Column{offset + 1}"
                        stem, number = name, 2
                        while name.casefold() in names:
                            name = f"{stem}{number}"
                            number += 1
                        names.add(name.casefold())
                        column.set("name", name)
                        set_header(self.root, ref, name)
                    self.warn(
                        "table_header_promoted",
                        table=table.get("name"),
                        message="Deleted table header; converted the surviving first row to unique text headers",
                    )
                if self.shift.point(bottom) is None and int(
                    table.get("totalsRowCount", "0")
                ):
                    table.set("totalsRowCount", "0")
                    table.set("totalsRowShown", "0")
            for node in list(table.iter()):
                if local(node) == "autoFilter":
                    self.autofilter(node)
                elif node is not table and node.get("ref"):
                    self.range_attribute(node, "ref")
        for node in list(self.root.iter(tag("tableParts"))):
            if len(node):
                node.set("count", str(len(node)))
            else:
                node.getparent().remove(node)

    def anchors_and_comments(self):
        for relation, part in self.p.relationships(self.part):
            kind = relation.get("Type", "").rsplit("/", 1)[-1].lower()
            if kind not in {
                "comments",
                "threadedcomment",
                "threadedcomments",
                "vmldrawing",
                "drawing",
            }:
                continue
            root = self.p.xml(part)
            if kind in {"comments", "threadedcomment", "threadedcomments"}:
                for node in list(root.iter()):
                    if node.get("ref"):
                        self.range_attribute(node, "ref")
            elif kind == "drawing":
                for marker in root.iter():
                    if marker.tag in {f"{{{DRAW}}}from", f"{{{DRAW}}}to"}:
                        node = marker.find(
                            f"{{{DRAW}}}{'row' if self.shift.axis == 'row' else 'col'}"
                        )
                        if node is not None:
                            node.text = str(
                                min(
                                    self.shift.limit - 1,
                                    self.shift.point(int(node.text) + 1, True) - 1,
                                )
                            )
            else:
                for shape in list(root.iter(f"{{{VML}}}shape")):
                    client = shape.find(f"{{{EXCEL}}}ClientData")
                    if client is None:
                        continue
                    axis = "Row" if self.shift.axis == "row" else "Column"
                    coord = client.find(f"{{{EXCEL}}}{axis}")
                    if coord is not None:
                        moved = self.shift.point(int(coord.text) + 1)
                        if moved is None and client.get("ObjectType") == "Note":
                            shape.getparent().remove(shape)
                            continue
                        coord.text = str(
                            min(self.shift.limit - 1, (moved or self.shift.at) - 1)
                        )
                    anchor = client.find(f"{{{EXCEL}}}Anchor")
                    if anchor is not None:
                        nums = [int(v.strip()) for v in anchor.text.split(",")]
                        original_nums = nums[:]
                        if len(nums) != 8:
                            raise ValueError(f"Invalid VML anchor in {part}")
                        for index in (2, 6) if self.shift.axis == "row" else (0, 4):
                            nums[index] = min(
                                self.shift.limit - 1,
                                self.shift.point(nums[index] + 1, True) - 1,
                            )
                        if nums != original_nums:
                            anchor.text = ", ".join(map(str, nums))

    def other_formulas(self):
        chart_parts = set()
        names = list(self.p.sheets)
        views = self.p.workbook.find(tag("bookViews"))
        active = (
            int(views[0].get("activeTab", "0"))
            if views is not None and len(views)
            else 0
        )
        default_sheet = names[min(active, len(names) - 1)]
        for part in self.p.original:
            if part in self.p.removed or not part.endswith(".xml"):
                continue
            root = self.p.xml(part)
            if local(root) == "pivotCacheDefinition":
                self.warn(
                    "pivot_cache_unchanged",
                    part=part,
                    message="Pivot cache source ranges were not adjusted",
                )
                continue
            context = self.owners.get(part)
            for node in root.iter():
                kind = local(node)
                if kind == "definedName" and part == "xl/workbook.xml":
                    scope = (
                        int(node.get("localSheetId"))
                        if node.get("localSheetId") is not None
                        else None
                    )
                    if scope is not None and not 0 <= scope < len(names):
                        raise ValueError("Invalid defined name localSheetId")
                    self.text_formula(
                        node,
                        names[scope] if scope is not None else default_sheet,
                        "names_rewritten",
                    )
                elif node.tag == f"{{{CHART}}}f":
                    old = node.text
                    if old:
                        node.text = self.shift.formula(old, context)
                        if node.text != old:
                            chart_parts.add(part)
                elif kind in FORMULA_TAGS or (kind == "f" and node.tag != tag("f")):
                    self.text_formula(node, context)
                elif kind == "hyperlink" and node.get("location"):
                    location = node.get("location")
                    prefix = (
                        "#"
                        if location.startswith("#") and not location.startswith("#REF!")
                        else ""
                    )
                    node.set(
                        "location",
                        prefix + self.shift.formula(location[len(prefix) :], context),
                    )
        self.report["charts_rewritten"] = len(chart_parts)

    def finalize(self):
        dimension = self.root.find(tag("dimension"))
        if dimension is not None:
            coords = [xy(c.get("r")) for c in self.root.iter(tag("c"))]
            if coords:
                a = address(min(c for c, r in coords), min(r for c, r in coords))
                b = address(max(c for c, r in coords), max(r for c, r in coords))
                dimension.set("ref", a if a == b else a + ":" + b)
            else:
                dimension.set("ref", "A1")
        for relation in list(self.p.xml("xl/_rels/workbook.xml.rels")):
            if relation.get("Type", "").endswith("/calcChain"):
                self.p.discard(resolve_part("xl/workbook.xml", relation.get("Target")))
                relation.getparent().remove(relation)
        self.p.discard("xl/calcChain.xml")
        calc = self.p.workbook.find(tag("calcPr"))
        if calc is None:
            calc = etree.Element(tag("calcPr"))
            after_calc = {
                "oleSize",
                "customWorkbookViews",
                "pivotCaches",
                "smartTagPr",
                "smartTagTypes",
                "webPublishing",
                "fileRecoveryPr",
                "webPublishObjects",
                "extLst",
            }
            next_node = next(
                (n for n in self.p.workbook if local(n) in after_calc), None
            )
            if next_node is not None:
                next_node.addprevious(calc)
            else:
                self.p.workbook.append(calc)
        calc.set("fullCalcOnLoad", "1")
        relpart = relation_path(self.part)
        if relpart in self.p.original:
            links = {n.get(f"{{{REL}}}id") for n in self.root.iter(tag("hyperlink"))}
            for rel in list(self.p.xml(relpart)):
                if (
                    rel.get("Type", "").endswith("/hyperlink")
                    and rel.get("Id") not in links
                ):
                    rel.getparent().remove(rel)

    def run(self):
        pending = []
        for sheet, part in self.p.sheets.items():
            root = self.p.xml(part)
            groups, arrays = self.rewrite_cells(root, sheet)
            pending.append((groups, arrays, root, sheet == self.sheet, sheet))
        self.shift_grid()
        for args in pending:
            self.finish_formulas(*args)
        self.sheet_ranges()
        self.tables()
        self.anchors_and_comments()
        self.other_formulas()
        self.finalize()


def formula_records(package):
    owners = package.owners()
    for part, data in package.original.items():
        if not part.endswith(".xml"):
            continue
        root = package.xml(part)
        sheet = owners.get(part)
        for cell, formula, text in cell_formulas(root):
            yield (
                {"part": part, "sheet": sheet, "cell": cell.get("r"), "kind": "cell"},
                text,
            )
            for attr in ("ref", "r1", "r2"):
                if "#REF!" in formula.get(attr, ""):
                    yield (
                        {
                            "part": part,
                            "sheet": sheet,
                            "cell": cell.get("r"),
                            "kind": "formula_range",
                        },
                        formula.get(attr),
                    )
        for node in root.iter():
            kind = local(node)
            if kind in FORMULA_TAGS | {"definedName"} or (
                kind == "f" and node.tag != tag("f")
            ):
                if node.text:
                    context = {"part": part, "sheet": sheet, "cell": None, "kind": kind}
                    if kind == "definedName":
                        context["name"] = node.get("name")
                        if node.get("localSheetId") is not None:
                            context["sheet"] = list(package.sheets)[
                                int(node.get("localSheetId"))
                            ]
                    parent = node.getparent()
                    while parent is not None:
                        if parent.get("sqref") or parent.get("ref"):
                            context["cell"] = parent.get("sqref") or parent.get("ref")
                            break
                        parent = parent.getparent()
                    yield context, node.text
            elif (
                node.tag == tag("c")
                and node.get("t") == "e"
                and node.find(tag("f")) is None
            ):
                if node.findtext(tag("v")) == "#REF!":
                    yield (
                        {
                            "part": part,
                            "sheet": sheet,
                            "cell": node.get("r"),
                            "kind": "error_cell",
                        },
                        "#REF!",
                    )
            elif kind == "hyperlink" and node.get("location"):
                location = node.get("location")
                if location.startswith("#") and not location.startswith("#REF!"):
                    location = location[1:]
                yield (
                    {
                        "part": part,
                        "sheet": sheet,
                        "cell": node.get("ref"),
                        "kind": "hyperlink",
                    },
                    location,
                )


def reference_warnings(package):
    warnings = []
    for location, text in formula_records(package):
        for index in range(error_count(text)):
            warnings.append(
                {
                    "type": "ref_error",
                    **location,
                    "occurrence": index + 1,
                    "formula": text,
                }
            )
        for token in formula_tokens(text):
            if (
                token.type == "OPERAND"
                and token.subtype == "RANGE"
                and "!" in token.value
            ):
                qualifier = token.value.rsplit("!", 1)[0]
                if "[" in qualifier:
                    warnings.append(
                        {
                            "type": "external_reference",
                            **location,
                            "reference": token.value,
                            "message": "External workbook reference left unchanged",
                        }
                    )
    return warnings


def self_check(source, original, removed, sheet, warnings):
    problems = []
    try:
        package = Package(source)
        if set(package.original) != set(original) - set(removed):
            problems.append(
                "Package part inventory differs from the intended inventory"
            )
        errors = 0
        for location, text in formula_records(package):
            try:
                validate_formula(text, location["kind"] == "definedName")
                errors += error_count(text)
            except Exception as exc:
                problems.append(f"Invalid formula/name at {location}: {exc}")
        if errors != sum(w.get("type") == "ref_error" for w in warnings):
            problems.append("#REF! occurrence count does not match warnings")
        for part, data in package.original.items():
            if not part.endswith((".xml", ".vml", ".rels")):
                continue
            for node in package.xml(part).iter():
                if not isinstance(node.tag, str):
                    continue
                for attr, value in node.attrib.items():
                    if etree.QName(attr).localname in {"ref", "sqref"}:
                        try:
                            if not value.strip():
                                raise ValueError("Empty range list")
                            for item in value.split():
                                if item != "#REF!":
                                    parse_range(item)
                        except ValueError as exc:
                            problems.append(
                                f"Invalid {local(node)}/@{attr} in {part}: {exc}"
                            )
                if local(node) == "sqref" and node.text:
                    for item in node.text.split():
                        try:
                            parse_range(item)
                        except ValueError as exc:
                            problems.append(f"Invalid sqref in {part}: {exc}")
        root = package.xml(package.sheets[sheet])
        data = root.find(tag("sheetData"))
        rows, seen = [], set()
        for row in data:
            number = int(row.get("r"))
            rows.append(number)
            if not 1 <= number <= MAX_ROW:
                problems.append(f"Row outside grid: {number}")
            columns = []
            for cell in row:
                if local(cell) != "c":
                    continue
                ref = cell.get("r")
                col, cell_row = xy(ref)
                if ref in seen:
                    problems.append(f"Duplicate cell: {ref}")
                seen.add(ref)
                columns.append(col)
                if cell_row != number:
                    problems.append(f"Cell {ref} is in row {number}")
            if columns != sorted(set(columns)):
                problems.append(f"Cells are unordered or duplicated in row {number}")
        if rows != sorted(set(rows)):
            problems.append("Rows are unordered or duplicated")
        for part in package.sheets.values():
            shared = {}
            for cell, formula, text in cell_formulas(package.xml(part)):
                if formula.get("t") == "shared":
                    si = formula.get("si")
                    shared.setdefault(si, []).append((cell, formula))
                if (
                    formula.get("t") in {"shared", "array"}
                    and formula.text
                    and not formula.get("ref")
                ):
                    problems.append(
                        f"Formula master missing range in {part}: {cell.get('r')}"
                    )
            for si, group in shared.items():
                masters = [(c, f) for c, f in group if f.text]
                if len(masters) != 1:
                    problems.append(
                        f"Shared group {si} in {part} must have exactly one master"
                    )
                elif masters[0][1].get("ref"):
                    kind, ends = parse_range(masters[0][1].get("ref"))
                    if kind != "cell":
                        problems.append(
                            f"Shared group {si} in {part} has a non-cell range"
                        )
                    else:
                        for cell, _ in group:
                            col, row = xy(cell.get("r"))
                            if not (
                                ends[0][1] <= col <= ends[-1][1]
                                and ends[0][2] <= row <= ends[-1][2]
                            ):
                                problems.append(
                                    f"Shared follower {cell.get('r')} lies outside group {si} in {part}"
                                )
    except Exception as exc:
        problems.append(f"{type(exc).__name__}: {exc}")
    return {"ok": False, "problems": problems} if problems else {"ok": True}


def position(value, axis):
    if axis == "row":
        if not re.fullmatch(r"[1-9][0-9]*", value):
            raise ValueError("--at and --copy-style-from must be positive row numbers")
        number = int(value)
        limit = MAX_ROW
    else:
        if not re.fullmatch(r"[A-Za-z]{1,3}", value):
            raise ValueError("--at and --copy-style-from must be column letters")
        number = column_index_from_string(value)
        limit = MAX_COL
    if number > limit:
        raise ValueError("Position is outside Excel's grid")
    return number


def parse_args(argv):
    if len(argv) < 2:
        raise ValueError(
            "Expected an operation and an input file; use --help for usage"
        )
    command, filename = argv[:2]
    operations = {
        "rows": "insert_rows",
        "columns": "insert_columns",
        "delete-rows": "delete_rows",
        "delete-columns": "delete_columns",
    }
    if command not in operations:
        raise ValueError(f"Unknown operation: {command}")
    flags, i = {}, 2
    while i < len(argv):
        flag = argv[i]
        if flag not in {
            "--sheet",
            "--at",
            "--count",
            "--copy-style-from",
            "--out",
            "--dry-run",
        }:
            raise ValueError(f"Unknown argument: {flag}")
        if flag in flags:
            raise ValueError(f"Duplicate argument: {flag}")
        if flag == "--dry-run":
            flags[flag] = True
            i += 1
        else:
            if i + 1 >= len(argv) or argv[i + 1].startswith("--"):
                raise ValueError(f"{flag} requires a value")
            flags[flag] = argv[i + 1]
            i += 2
    if "--sheet" not in flags or "--at" not in flags:
        raise ValueError("--sheet and --at are required")
    axis = "row" if command.endswith("rows") else "column"
    at = position(flags["--at"], axis)
    try:
        count = int(flags.get("--count", "1"))
    except ValueError:
        raise ValueError("--count must be an integer") from None
    if count < 1:
        raise ValueError("--count must be at least 1")
    if at + count - 1 > (MAX_ROW if axis == "row" else MAX_COL):
        raise ValueError("The requested band extends outside Excel's grid")
    copy_from = (
        position(flags["--copy-style-from"], axis)
        if "--copy-style-from" in flags
        else None
    )
    delete = command.startswith("delete-")
    if delete and copy_from is not None:
        raise ValueError("--copy-style-from is only valid for insertion")
    return filename, operations[command], flags, axis, at, count, delete, copy_from


def execute(argv):
    filename, operation, flags, axis, at, count, delete, copy_from = parse_args(argv)
    source = Path(filename).expanduser().resolve()
    destination = Path(flags.get("--out", str(source))).expanduser().resolve()
    package = Package(source)
    sheet = flags["--sheet"]
    if sheet not in package.sheets:
        raise ValueError(f"Unknown sheet: {sheet}")
    report = {
        "status": "success",
        "sheet": sheet,
        "operation": operation,
        "at": flags["--at"],
        "count": count,
        "cells_shifted": 0,
        "formulas_rewritten": 0,
        "names_rewritten": 0,
        "charts_rewritten": 0,
        "parts_changed": [],
        "warnings": [],
    }
    if flags.get("--dry-run"):
        report["dry_run"] = True
    Editor(package, sheet, axis, at, count, delete, copy_from, report).run()
    payload = package.payload()
    report["parts_changed"] = sorted(
        p for p in package.original if payload.get(p) != package.original[p]
    )
    # A dry run does not even create a temporary file. Both paths reopen a ZIP
    # serialized by precisely the same writer before checking or committing it.
    buffer = io.BytesIO()
    package.write(buffer, payload)
    buffer.seek(0)
    computed = Package(buffer)
    report["warnings"].extend(reference_warnings(computed))
    if flags.get("--dry-run"):
        buffer.seek(0)
        report["self_check"] = self_check(
            buffer, package.original, package.removed, sheet, report["warnings"]
        )
    else:
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=".insert-", suffix=".xlsx", dir=destination.parent, delete=False
            ) as handle:
                temporary = Path(handle.name)
                handle.write(buffer.getvalue())
                handle.flush()
                os.fsync(handle.fileno())
            os.chmod(temporary, source.stat().st_mode & 0o777)
            report["self_check"] = self_check(
                temporary, package.original, package.removed, sheet, report["warnings"]
            )
            if report["self_check"]["ok"]:
                os.replace(temporary, destination)
        finally:
            if temporary is not None and temporary.exists():
                temporary.unlink()
    if not report["self_check"]["ok"]:
        report["status"] = "error"
        report["message"] = "Self-check failed; no output was committed"
    return report


def main(argv):
    if "--help" in argv or "-h" in argv:
        print(__doc__.strip())
        return 0
    try:
        report = execute(argv)
    except Exception as exc:
        report = {"status": "error", "message": f"{type(exc).__name__}: {exc}"}
    print(json.dumps(report, indent=2))
    return 1 if report["status"] == "error" else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
