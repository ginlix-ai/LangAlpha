"""Locks the xlsx skill's insert.py: every reference class Excel repoints, on generated workbooks."""

import copy
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import struct
import zlib
import zipfile

from lxml import etree
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[4]
SCRIPT = (
    ROOT
    / "plugins"
    / "langalpha_deliverables"
    / "skills"
    / "xlsx"
    / "scripts"
    / "insert.py"
)
spec = importlib.util.spec_from_file_location("insert_script", SCRIPT)
insert = importlib.util.module_from_spec(spec)
spec.loader.exec_module(insert)
S = insert.SML
R = insert.REL
C = insert.CHART
D = insert.DRAW
P = "http://schemas.openxmlformats.org/package/2006/relationships"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
TC = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
NS = {
    "s": S,
    "c": C,
    "xdr": D,
    "x": insert.EXCEL,
    "v": insert.VML,
    "x14": X14,
    "xm": XM,
}


@pytest.fixture(autouse=True)
def library_temporary_files_stay_here(monkeypatch):
    monkeypatch.setattr(tempfile, "tempdir", str(ROOT))


@pytest.fixture
def tmp_path():
    # Keep fixtures and all test output inside the task directory.
    with tempfile.TemporaryDirectory(prefix=".insert-test-", dir=ROOT) as path:
        yield Path(path)


def payload(path):
    with zipfile.ZipFile(path) as archive:
        return {p: archive.read(p) for p in archive.namelist()}


def patch_package(path, change):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        data = {p: archive.read(p) for p in archive.namelist()}
        comment = archive.comment
    change(data)
    with zipfile.ZipFile(path, "w") as archive:
        archive.comment = comment
        for info in infos:
            if info.filename in data:
                archive.writestr(info, data.pop(info.filename))
        for name, blob in data.items():
            archive.writestr(name, blob)


def xml(data, part="xl/worksheets/sheet2.xml"):
    return etree.fromstring(data[part])


def save_xml(data, part, root):
    data[part] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)


def formula(root, ref):
    return root.findtext(f"s:sheetData/s:row/s:c[@r='{ref}']/s:f", namespaces=NS)


def cell_node(root, ref):
    return root.find(f"s:sheetData/s:row/s:c[@r='{ref}']", namespaces=NS)


def run(path, operation="rows", at="4", *extra):
    args = [operation, str(path), "--sheet", "X", "--at", at, *map(str, extra)]
    return insert.execute(args)


def assert_preserved(before, after, report):
    changed = set(report["parts_changed"])
    assert changed == {part for part in before if before[part] != after.get(part)}
    for part in before.keys() - changed:
        assert after[part] == before[part], part
    assert before.keys() - after.keys() <= changed


def basic(path, rows=12, columns=8):
    wb = Workbook()
    ws = wb.active
    ws.title = "X"
    for row in range(1, rows + 1):
        for col in range(1, columns + 1):
            ws.cell(row, col, row * 100 + col)
    wb.create_sheet("Other")
    wb.save(path)
    return path


def rich(path):
    wb = Workbook()
    wb.active.title = "Start"
    ws = wb.create_sheet("X")
    last = wb.create_sheet("End")
    other = wb.create_sheet("Other Sheet")
    untouched = wb.create_sheet("Untouched")
    untouched["A1"] = "This part must stay byte-identical"
    untouched["B2"] = "=1+2"
    untouched["C3"].comment = Comment("Leave this comment alone", "Test")
    for row in range(1, 13):
        for col in range(1, 15):
            ws.cell(row, col, row * 100 + col)
    for offset, name in enumerate(["Label", "Value", "Cost", "Profit"], 2):
        ws.cell(2, offset, name)
    ws.add_table(Table(displayName="Sales", ref="B2:E8"))
    ws["M2"] = (
        '=SUM($B$3:$D$8)+SUM(3:8)+SUM(B:D)+LOG10(B3)+ATAN2(B3,C3)+IF(B3>0,"B3",0)+Sales[Cost]+Stable'
    )
    ws["M3"] = "=D4"
    ws["M12"] = "='Other Sheet'!B4+Start!B4+SUM(Start:End!B3:D8)"
    other["A1"] = "=X!$B$3+SUM('X'!B3:D8)+SUM(Start:End!B3:D8)+SUM(X!B:D)+SUM(X!3:8)"
    other["A2"] = "=B4+LOG10(B3)+ATAN2(B3,C3)+Sales[Cost]+Stable"
    last["A1"] = "='Other Sheet'!B4"
    ws["D5"].comment = Comment("A surviving or deleted comment", "Tester")
    ws["F6"].hyperlink = "https://example.invalid/kept"
    ws.merge_cells("B10:D11")
    ws.conditional_formatting.add(
        "B3:D8",
        FormulaRule(formula=["B3>0"], fill=PatternFill("solid", fgColor="FF0000")),
    )
    dv = DataValidation(type="list", formula1="$B$3:$B$8")
    dv.add("C3:C8")
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="whole", operator="between", formula1="B3", formula2="D8")
    dv2.add("F3:F8")
    ws.add_data_validation(dv2)
    ws.auto_filter.ref = "B2:F9"
    ws.auto_filter.add_filter_column(3, ["3"])
    ws.auto_filter.add_sort_condition("D3:D9")
    ws.freeze_panes = "D4"
    ws.sheet_view.selection[-1].activeCell = "D5"
    ws.sheet_view.selection[-1].sqref = "D5 F6:G8"
    ws.print_area = "B2:F9"
    ws.print_title_rows = "2:5"
    ws.print_title_cols = "B:D"
    ws.row_dimensions[5].hidden = True
    ws.row_dimensions[5].outlineLevel = 2
    ws.row_dimensions[5].height = 29
    ws.column_dimensions.group("D", "F", hidden=True)
    ws.column_dimensions["D"].width = 23
    ws["D5"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["E5"].number_format = "0.00%"
    wb.defined_names.add(DefinedName("Stable", attr_text="'X'!$B$3:$D$8"))
    wb.defined_names.add(DefinedName("Constant", attr_text='"B3"'))
    ws.defined_names.add(DefinedName("LocalBand", attr_text="$B$3:$D$8"))
    other.defined_names.add(DefinedName("OtherBand", attr_text="$B$3:$D$8"))
    chart = BarChart()
    chart.add_data(
        Reference(ws, min_col=3, max_col=4, min_row=2, max_row=8), titles_from_data=True
    )
    chart.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=7, row=2), to=AnchorMarker(col=11, row=9)
    )
    ws.add_chart(chart)
    other_chart = BarChart()
    other_chart.add_data(
        Reference(ws, min_col=3, min_row=2, max_row=8), titles_from_data=True
    )
    other.add_chart(other_chart, "H3")
    wb.save(path)

    def patch(data):
        part = "xl/worksheets/sheet2.xml"
        root = xml(data, part)
        for row in root.findall("s:sheetData/s:row", NS):
            row.set("spans", "1:14")
        for row in range(3, 9):
            cell = cell_node(root, f"J{row}")
            cell[:] = []
            f = etree.SubElement(cell, f"{{{S}}}f", t="shared", si="0")
            if row == 3:
                f.set("ref", "J3:J8")
                f.text = "B3*2"
            etree.SubElement(cell, f"{{{S}}}v").text = "0"
        cell = cell_node(root, "K3")
        cell[:] = []
        etree.SubElement(cell, f"{{{S}}}f", t="array", ref="K3:K8").text = "B3:B8*2"
        extlist = etree.SubElement(root, f"{{{S}}}extLst")
        ext = etree.SubElement(extlist, f"{{{S}}}ext", uri="{spark-test}")
        groups = etree.SubElement(
            ext, f"{{{X14}}}sparklineGroups", nsmap={"x14": X14, "xm": XM}
        )
        group = etree.SubElement(groups, f"{{{X14}}}sparklineGroup")
        etree.SubElement(group, f"{{{XM}}}f").text = "X!B3:B8"
        sparks = etree.SubElement(group, f"{{{X14}}}sparklines")
        spark = etree.SubElement(sparks, f"{{{X14}}}sparkline")
        etree.SubElement(spark, f"{{{XM}}}f").text = "X!B3:D3"
        etree.SubElement(spark, f"{{{XM}}}sqref").text = "L5"
        save_xml(data, part, root)
        relpart = "xl/worksheets/_rels/sheet2.xml.rels"
        rels = xml(data, relpart)
        etree.SubElement(
            rels,
            f"{{{P}}}Relationship",
            Id="rThread",
            Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment",
            Target="../threadedComments/threadedComment1.xml",
        )
        save_xml(data, relpart, rels)
        thread = etree.Element(f"{{{TC}}}ThreadedComments", nsmap={None: TC})
        tc = etree.SubElement(
            thread, f"{{{TC}}}threadedComment", ref="E5", personId="person", id="thread"
        )
        etree.SubElement(tc, f"{{{TC}}}text").text = "Thread text"
        data["xl/threadedComments/threadedComment1.xml"] = etree.tostring(thread)
        vml_part = next(p for p in data if p.endswith("commentsDrawing1.vml"))
        vml = xml(data, vml_part)
        client = vml.find(".//x:ClientData", NS)
        etree.SubElement(
            client, f"{{{insert.EXCEL}}}Anchor"
        ).text = "3, 15, 4, 2, 6, 20, 8, 4"
        save_xml(data, vml_part, vml)
        chain = etree.Element(f"{{{S}}}calcChain", nsmap={None: S})
        etree.SubElement(chain, f"{{{S}}}c", r="M3", i="2")
        data["xl/calcChain.xml"] = etree.tostring(chain)
        rels = xml(data, "xl/_rels/workbook.xml.rels")
        etree.SubElement(
            rels,
            f"{{{P}}}Relationship",
            Id="rChain",
            Type=R + "/calcChain",
            Target="calcChain.xml",
        )
        save_xml(data, "xl/_rels/workbook.xml.rels", rels)
        types = xml(data, "[Content_Types].xml")
        for partname, contenttype in [
            (
                "/xl/calcChain.xml",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
            ),
            (
                "/xl/threadedComments/threadedComment1.xml",
                "application/vnd.ms-excel.threadedcomments+xml",
            ),
        ]:
            etree.SubElement(
                types,
                "{" + types.nsmap[None] + "}Override",
                PartName=partname,
                ContentType=contenttype,
            )
        save_xml(data, "[Content_Types].xml", types)
        data["custom/opaque.bin"] = b"\x00\xffUNTOUCHED\x00"

        # Generate a real single-pixel PNG and an image anchor without Pillow.
        def chunk(kind, blob):
            return (
                struct.pack(">I", len(blob))
                + kind
                + blob
                + struct.pack(">I", zlib.crc32(kind + blob))
            )

        data["xl/media/pixel.png"] = (
            b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
            + chunk(b"IEND", b"")
        )
        drawing = xml(data, "xl/drawings/drawing1.xml")
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        anchor = etree.SubElement(drawing, f"{{{D}}}oneCellAnchor")
        start = etree.SubElement(anchor, f"{{{D}}}from")
        for key, value in [
            ("col", "7"),
            ("colOff", "0"),
            ("row", "4"),
            ("rowOff", "0"),
        ]:
            etree.SubElement(start, f"{{{D}}}{key}").text = value
        etree.SubElement(anchor, f"{{{D}}}ext", cx="9525", cy="9525")
        pic = etree.SubElement(anchor, f"{{{D}}}pic")
        nv = etree.SubElement(pic, f"{{{D}}}nvPicPr")
        etree.SubElement(nv, f"{{{D}}}cNvPr", id="2", name="Pixel")
        etree.SubElement(nv, f"{{{D}}}cNvPicPr")
        fill = etree.SubElement(pic, f"{{{D}}}blipFill")
        etree.SubElement(fill, f"{{{a}}}blip", {f"{{{R}}}embed": "rPixel"})
        etree.SubElement(etree.SubElement(fill, f"{{{a}}}stretch"), f"{{{a}}}fillRect")
        geometry = etree.SubElement(
            etree.SubElement(pic, f"{{{D}}}spPr"), f"{{{a}}}prstGeom", prst="rect"
        )
        etree.SubElement(geometry, f"{{{a}}}avLst")
        etree.SubElement(anchor, f"{{{D}}}clientData")
        save_xml(data, "xl/drawings/drawing1.xml", drawing)
        rels = xml(data, "xl/drawings/_rels/drawing1.xml.rels")
        etree.SubElement(
            rels,
            f"{{{P}}}Relationship",
            Id="rPixel",
            Type=R + "/image",
            Target="../media/pixel.png",
        )
        save_xml(data, "xl/drawings/_rels/drawing1.xml.rels", rels)
        types = xml(data, "[Content_Types].xml")
        etree.SubElement(
            types,
            "{" + types.nsmap[None] + "}Default",
            Extension="png",
            ContentType="image/png",
        )
        save_xml(data, "[Content_Types].xml", types)

    patch_package(path, patch)
    return path


@pytest.mark.parametrize(
    "op,at,expected",
    [
        (
            "rows",
            "4",
            dict(
                band="B3:D9",
                rowband="3:9",
                colband="B:D",
                merge="B11:D12",
                table="B2:E9",
                filter="B2:F10",
                fid="3",
                freeze="D5",
                xs="3",
                ys="4",
                comment="D6",
                thread="E6",
                hyperlink="F7",
                spark="L6",
                shared="J3:J9",
                array="K3:K9",
                end="D9",
                dim="A1:N13",
                outline_row="6",
                columns=[("4", "6")],
                from_rc=(2, 7),
                to_rc=(10, 11),
                vml=[3, 15, 5, 2, 6, 20, 9, 4],
                cross="X!$B$3+SUM('X'!B3:D9)+SUM(Start:End!B3:D9)+SUM(X!B:D)+SUM(X!3:9)",
                single="D5",
                m12="M13",
            ),
        ),
        (
            "delete-rows",
            "4",
            dict(
                band="B3:D7",
                rowband="3:7",
                colband="B:D",
                merge="B9:D10",
                table="B2:E7",
                filter="B2:F8",
                fid="3",
                freeze="D4",
                xs="3",
                ys="3",
                comment="D4",
                thread="E4",
                hyperlink="F5",
                spark="L4",
                shared="J3:J7",
                array="K3:K7",
                end="D7",
                dim="A1:N11",
                outline_row="4",
                columns=[("4", "6")],
                from_rc=(2, 7),
                to_rc=(8, 11),
                vml=[3, 15, 3, 2, 6, 20, 7, 4],
                cross="X!$B$3+SUM('X'!B3:D7)+SUM(Start:End!B3:D7)+SUM(X!B:D)+SUM(X!3:7)",
                single="#REF!",
                m12="M11",
            ),
        ),
        (
            "columns",
            "D",
            dict(
                band="B3:E8",
                rowband="3:8",
                colband="B:E",
                merge="B10:E11",
                table="B2:F8",
                filter="B2:G9",
                fid="4",
                freeze="E4",
                xs="4",
                ys="3",
                comment="E5",
                thread="F5",
                hyperlink="G6",
                spark="M5",
                shared="K3:K8",
                array="L3:L8",
                end="E8",
                dim="A1:O12",
                outline_row="5",
                columns=[("5", "7")],
                from_rc=(2, 8),
                to_rc=(9, 12),
                vml=[4, 15, 4, 2, 7, 20, 8, 4],
                cross="X!$B$3+SUM('X'!B3:E8)+SUM(Start:End!B3:E8)+SUM(X!B:E)+SUM(X!3:8)",
                single="E4",
                m12="N12",
            ),
        ),
        (
            "delete-columns",
            "D",
            dict(
                band="B3:C8",
                rowband="3:8",
                colband="B:C",
                merge="B10:C11",
                table="B2:D8",
                filter="B2:E9",
                fid="2",
                freeze="D4",
                xs="3",
                ys="3",
                comment=None,
                thread="D5",
                hyperlink="E6",
                spark="K5",
                shared="I3:I8",
                array="J3:J8",
                end="C8",
                dim="A1:M12",
                outline_row="5",
                columns=[("4", "5")],
                from_rc=(2, 6),
                to_rc=(9, 10),
                vml=None,
                cross="X!$B$3+SUM('X'!B3:C8)+SUM(Start:End!B3:C8)+SUM(X!B:C)+SUM(X!3:8)",
                single="#REF!",
                m12="L12",
            ),
        ),
    ],
)
def test_comprehensive_four_operations(tmp_path, op, at, expected):
    path = rich(tmp_path / "rich.xlsx")
    before = payload(path)
    report = run(path, op, at)
    assert report["status"] == "success", report
    assert report["self_check"] == {"ok": True}
    assert report["at"] == at and report["count"] == 1
    assert (
        report["operation"]
        == {
            "rows": "insert_rows",
            "columns": "insert_columns",
            "delete-rows": "delete_rows",
            "delete-columns": "delete_columns",
        }[op]
    )
    after = payload(path)
    assert_preserved(before, after, report)
    assert after["xl/worksheets/sheet5.xml"] == before["xl/worksheets/sheet5.xml"]
    assert after["xl/styles.xml"] == before["xl/styles.xml"]
    assert after["custom/opaque.bin"] == before["custom/opaque.bin"]
    assert after["xl/drawings/drawing2.xml"] == before["xl/drawings/drawing2.xml"]
    root = xml(after)
    assert root.find("s:dimension", NS).get("ref") == expected["dim"]
    assert root.find("s:mergeCells/s:mergeCell", NS).get("ref") == expected["merge"]
    cf = root.find("s:conditionalFormatting", NS)
    assert cf.get("sqref") == expected["band"]
    assert cf.findtext("s:cfRule/s:formula", namespaces=NS) == "B3>0"
    dv = root.find("s:dataValidations/s:dataValidation", NS)
    assert (
        dv.findtext("s:formula1", namespaces=NS)
        == "$B$3:$B$" + expected["band"].split(":")[1][1:]
    )
    dv2 = root.findall("s:dataValidations/s:dataValidation", NS)[1]
    assert dv2.findtext("s:formula2", namespaces=NS) == (
        "#REF!" if op == "delete-columns" else expected["end"]
    )
    auto = root.find("s:autoFilter", NS)
    assert auto.get("ref") == expected["filter"]
    assert auto.find("s:filterColumn", NS).get("colId") == expected["fid"]
    pane = root.find("s:sheetViews/s:sheetView/s:pane", NS)
    assert pane.get("topLeftCell") == expected["freeze"]
    assert pane.get("xSplit") == expected["xs"]
    assert pane.get("ySplit") == expected["ys"]
    selections = root.findall("s:sheetViews/s:sheetView/s:selection", NS)
    assert [s.get("pane") for s in selections] == [
        "topRight",
        "bottomLeft",
        "bottomRight",
    ]
    expected_selection = {
        "rows": ("D6", "D6 F7:G9"),
        "delete-rows": ("D4", "D4 F5:G7"),
        "columns": ("E5", "E5 G6:H8"),
        "delete-columns": ("E6", "E6:F8"),
    }[op]
    assert (
        selections[-1].get("activeCell"),
        selections[-1].get("sqref"),
    ) == expected_selection
    assert root.find("s:hyperlinks/s:hyperlink", NS).get("ref") == expected["hyperlink"]
    row = root.find(f"s:sheetData/s:row[@r='{expected['outline_row']}']", NS)
    assert (
        row.get("hidden") == "1"
        and row.get("outlineLevel") == "2"
        and row.get("ht") == "29"
    )
    assert row.get("spans") == (
        "1:15" if op == "columns" else "1:13" if op == "delete-columns" else "1:14"
    )
    cols = root.findall("s:cols/s:col", NS)
    assert [(c.get("min"), c.get("max")) for c in cols] == expected["columns"]
    assert all(
        c.get("hidden") == "1"
        and c.get("outlineLevel") == "1"
        and c.get("width") == "23"
        for c in cols
    )
    assert root.find(".//x14:sparkline/xm:sqref", NS).text == expected["spark"]
    assert (
        root.find(".//x14:sparklineGroup/xm:f", NS).text
        == "X!B3:B" + expected["band"].split(":")[1][1:]
    )
    shared = root.find(".//s:f[@t='shared'][@ref]", NS)
    assert shared.get("ref") == expected["shared"]
    assert shared.get("si") == "0" and shared.text == "B3*2"
    array = root.find(".//s:f[@t='array']", NS)
    assert array.get("ref") == expected["array"]
    assert array.text == "B3:B" + expected["band"].split(":")[1][1:] + "*2"
    master_ref = {
        "rows": "M2",
        "delete-rows": "M2",
        "columns": "N2",
        "delete-columns": "L2",
    }[op]
    primary = formula(root, master_ref)
    finish = expected["band"].split(":")[1]
    abs_band = "$B$3:$" + finish[0] + "$" + finish[1:]
    assert (
        primary
        == f'SUM({abs_band})+SUM({expected["rowband"]})+SUM({expected["colband"]})+LOG10(B3)+ATAN2(B3,C3)+IF(B3>0,"B3",0)+Sales[Cost]+Stable'
    )
    assert formula(root, master_ref[:-1] + "3") == expected["single"]
    assert (
        formula(root, expected["m12"])
        == "'Other Sheet'!B4+Start!B4+SUM(Start:End!" + expected["band"] + ")"
    )
    other = xml(after, "xl/worksheets/sheet4.xml")
    assert formula(other, "A1") == expected["cross"]
    assert formula(other, "A2") == formula(
        xml(before, "xl/worksheets/sheet4.xml"), "A2"
    )
    names = {
        n.get("name"): n.text
        for n in xml(after, "xl/workbook.xml").findall(
            "s:definedNames/s:definedName", NS
        )
    }
    assert names["Stable"] == "'X'!" + abs_band
    assert names["LocalBand"] == abs_band
    assert names["OtherBand"] == "$B$3:$D$8"
    assert names["Constant"] == '"B3"'
    print_areas = {
        "rows": "'X'!$B$2:$F$10",
        "delete-rows": "'X'!$B$2:$F$8",
        "columns": "'X'!$B$2:$G$9",
        "delete-columns": "'X'!$B$2:$E$9",
    }
    assert names["_xlnm.Print_Area"] == print_areas[op]
    print_titles = {
        "rows": "'X'!$2:$6,'X'!$B:$D",
        "delete-rows": "'X'!$2:$4,'X'!$B:$D",
        "columns": "'X'!$2:$5,'X'!$B:$E",
        "delete-columns": "'X'!$2:$5,'X'!$B:$C",
    }
    assert names["_xlnm.Print_Titles"] == print_titles[op]
    assert "xl/calcChain.xml" not in after
    assert b"calcChain" not in after["xl/_rels/workbook.xml.rels"]
    assert b"calcChain" not in after["[Content_Types].xml"]
    assert (
        xml(after, "xl/workbook.xml").find("s:calcPr", NS).get("fullCalcOnLoad") == "1"
    )
    table = xml(after, "xl/tables/table1.xml")
    assert table.get("ref") == expected["table"]
    assert table.find("s:autoFilter", NS).get("ref") == expected["table"]
    table_cols = table.findall("s:tableColumns/s:tableColumn", NS)
    expected_names = (
        ["Label", "Value", "Column1", "Cost", "Profit"]
        if op == "columns"
        else ["Label", "Value", "Profit"]
        if op == "delete-columns"
        else ["Label", "Value", "Cost", "Profit"]
    )
    assert [c.get("name") for c in table_cols] == expected_names
    assert table.find("s:tableColumns", NS).get("count") == str(len(expected_names))
    if op == "columns":
        assert cell_node(root, "D2").findtext("s:is/s:t", namespaces=NS) == "Column1"
        assert any(w["type"] == "table_header" for w in report["warnings"])
    chart = xml(after, "xl/charts/chart1.xml")
    refs = chart.findall(".//c:f", NS)
    row_end = expected["band"].split(":")[1][1:]
    assert refs[1].text == "'X'!$C$3:$C$" + row_end
    assert report["charts_rewritten"] == (2 if "rows" in op else 1)
    drawing = xml(after, "xl/drawings/drawing1.xml")
    for marker, pair in [("from", expected["from_rc"]), ("to", expected["to_rc"])]:
        assert (
            int(drawing.findtext(f".//xdr:{marker}/xdr:row", namespaces=NS)),
            int(drawing.findtext(f".//xdr:{marker}/xdr:col", namespaces=NS)),
        ) == pair
    picture = drawing.find("xdr:oneCellAnchor", NS)
    assert (
        int(picture.findtext("xdr:from/xdr:row", namespaces=NS)),
        int(picture.findtext("xdr:from/xdr:col", namespaces=NS)),
    ) == {
        "rows": (5, 7),
        "delete-rows": (3, 7),
        "columns": (4, 8),
        "delete-columns": (4, 6),
    }[op]
    assert before["xl/media/pixel.png"] == after["xl/media/pixel.png"]
    comments = xml(after, "xl/comments/comment1.xml").findall(
        "s:commentList/s:comment", NS
    )
    assert [c.get("ref") for c in comments] == (
        [expected["comment"]] if expected["comment"] else []
    )
    thread = xml(after, "xl/threadedComments/threadedComment1.xml")
    assert thread[0].get("ref") == expected["thread"]
    vml = xml(after, "xl/drawings/commentsDrawing1.vml")
    anchors = vml.findall(".//x:Anchor", NS)
    assert (
        [int(v) for v in anchors[0].text.split(",")] if anchors else None
    ) == expected["vml"]
    ref_warnings = [w for w in report["warnings"] if w["type"] == "ref_error"]
    if op.startswith("delete"):
        assert any(
            w["cell"] == master_ref[:-1] + "3" and w["sheet"] == "X"
            for w in ref_warnings
        )
    else:
        assert not ref_warnings


@pytest.mark.parametrize(
    "axis,delete,at,count,original,expected",
    [
        ("row", False, 3, 2, "A3:B5", "A5:B7"),
        ("row", False, 4, 2, "$A$3:B$5", "$A$3:B$7"),
        ("row", False, 6, 2, "A3:B5", "A3:B5"),
        ("column", False, 2, 2, "B3:D5", "D3:F5"),
        ("column", False, 3, 2, "$B3:D$5", "$B3:F$5"),
        ("column", False, 5, 2, "B3:D5", "B3:D5"),
        ("row", True, 4, 2, "A3:B8", "A3:B6"),
        ("row", True, 3, 2, "A3:B8", "A3:B6"),
        ("row", True, 7, 2, "A3:B8", "A3:B6"),
        ("row", True, 2, 8, "A3:B8", None),
        ("column", True, 3, 2, "B3:F8", "B3:D8"),
        ("column", True, 2, 2, "B3:F8", "B3:D8"),
        ("column", True, 5, 2, "B3:F8", "B3:D8"),
        ("column", True, 2, 8, "B3:F8", None),
        ("row", False, 4, 2, "$3:8", "$3:10"),
        ("column", True, 3, 2, "$B:F", "$B:D"),
        ("row", True, 1, 9, "B:F", "B:F"),
        ("column", True, 1, 9, "3:8", "3:8"),
        ("row", False, 4, 2, "A8:B3", "A10:B3"),
        ("row", True, 4, 1, "$A$4", None),
        ("column", True, 4, 1, "$D4", None),
    ],
)
def test_reference_boundaries(axis, delete, at, count, original, expected):
    shift = insert.Shift("X", ["Start", "X", "End"], axis, at, count, delete)
    assert shift.range(original) == expected


@pytest.mark.parametrize(
    "axis,at,expected",
    [
        (
            "row",
            3,
            "SUM('X'!$B4, X!B$5, Start:End!B4, 'Start:End'!B4, Other!B3, [1]X!B3, '[book.xlsx]X'!B3, Sales[Cost], Stable, LOG10(B4), ATAN2(B4,C4), \"B3 #REF!\", #N/A)  +  B4",
        ),
        (
            "column",
            2,
            "SUM('X'!$C3, X!C$4, Start:End!C3, 'Start:End'!C3, Other!B3, [1]X!B3, '[book.xlsx]X'!B3, Sales[Cost], Stable, LOG10(C3), ATAN2(C3,D3), \"B3 #REF!\", #N/A)  +  C3",
        ),
    ],
)
def test_tokenizer_protects_non_references(axis, at, expected):
    text = "SUM('X'!$B3, X!B$4, Start:End!B3, 'Start:End'!B3, Other!B3, [1]X!B3, '[book.xlsx]X'!B3, Sales[Cost], Stable, LOG10(B3), ATAN2(B3,C3), \"B3 #REF!\", #N/A)  +  B3"
    shift = insert.Shift("X", ["Start", "X", "End", "Other"], axis, at, 1, False)
    assert shift.formula(text, "X") == expected
    assert (
        shift.formula("B3+Stable+LOG10(B3)+Other!B3", "Other")
        == "B3+Stable+LOG10(B3)+Other!B3"
    )
    assert shift.formula("SUM(End:Other!B3)", "X") == "SUM(End:Other!B3)"
    assert shift.formula("SUM(End:Start!B3)", "Other") != "SUM(End:Start!B3)"


def test_quoted_sheet_escaped_apostrophe_and_3d():
    shift = insert.Shift(
        "O'Brien", ["First Sheet", "O'Brien", "Last Sheet"], "row", 3, 1, False
    )
    text = "'O''Brien'!A3+'First Sheet:Last Sheet'!A3+'First Sheet':'Last Sheet'!A3"
    assert (
        shift.formula(text, "First Sheet")
        == "'O''Brien'!A4+'First Sheet:Last Sheet'!A4+'First Sheet':'Last Sheet'!A4"
    )
    apostrophe = insert.Shift("Bob'", ["Bob'"], "row", 2, 1, False)
    assert apostrophe.formula("'Bob'''!A2", "Other") == "'Bob'''!A3"


@pytest.mark.parametrize(
    "operation,at,source,new_refs,old_ref,new_old_ref",
    [
        ("rows", "4", "5", ["D4", "E4", "D5", "E5"], "D5", "D7"),
        ("columns", "D", "D", ["D5", "E5"], "D5", "F5"),
    ],
)
def test_copy_original_styles_and_dimensions(
    tmp_path, operation, at, source, new_refs, old_ref, new_old_ref
):
    path = basic(tmp_path / "style.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws["D5"].fill = PatternFill("solid", fgColor="FF0000")
    ws["E5"].number_format = "0.00%"
    ws.row_dimensions[5].height = 37
    ws.row_dimensions[5].hidden = True
    ws.column_dimensions["D"].width = 29
    ws.column_dimensions["D"].hidden = True
    ws["D5"] = "=A5"
    wb.save(path)
    before = payload(path)
    report = run(path, operation, at, "--count", "2", "--copy-style-from", source)
    assert report["self_check"]["ok"]
    out = load_workbook(path)
    for ref in new_refs:
        assert out["X"][ref].value is None
        assert out["X"][ref].has_style
    original_cell = cell_node(xml(before, "xl/worksheets/sheet1.xml"), old_ref)
    changed = xml(payload(path), "xl/worksheets/sheet1.xml")
    assert cell_node(changed, new_old_ref).get("s") == original_cell.get("s")
    assert cell_node(changed, new_refs[0]).get("s") == original_cell.get("s")
    if operation == "rows":
        assert out["X"].row_dimensions[4].height == 37
        assert out["X"].row_dimensions[5].height == 37
        assert not out["X"].row_dimensions[4].hidden
        assert out["X"].row_dimensions[7].hidden
        assert out["X"]["D7"].value == "=A7"
    else:
        assert out["X"].column_dimensions["D"].width == 29
        assert out["X"].column_dimensions["E"].width == 29
        assert not out["X"].column_dimensions["D"].hidden
        assert out["X"].column_dimensions["F"].hidden


@pytest.mark.parametrize("operation,at", [("rows", "5"), ("columns", "D")])
def test_no_style_source_leaves_inserted_band_unstyled(tmp_path, operation, at):
    path = basic(tmp_path / "blank.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws["D5"].fill = PatternFill("solid", fgColor="FF0000")
    ws.column_dimensions.group("B", "F", hidden=True)
    wb.save(path)
    assert run(path, operation, at)["status"] == "success"
    out = load_workbook(path)["X"]
    if operation == "rows":
        assert all(c.value is None and not c.has_style for c in out[5])
    else:
        assert all(c.value is None and not c.has_style for c in out["D"])
        assert not out.column_dimensions["D"].hidden
        root = xml(payload(path), "xl/worksheets/sheet1.xml")
        assert [
            (c.get("min"), c.get("max")) for c in root.findall("s:cols/s:col", NS)
        ] == [("2", "3"), ("5", "7")]


def test_dry_run_and_out_atomicity(tmp_path, monkeypatch):
    path = basic(tmp_path / "source.xlsx")
    source = path.read_bytes()
    output = tmp_path / "out.xlsx"
    entries = sorted(tmp_path.iterdir())
    dry = run(path, "rows", "3", "--out", output, "--dry-run")
    assert dry["dry_run"] is True and dry["self_check"] == {"ok": True}
    assert path.read_bytes() == source and sorted(tmp_path.iterdir()) == entries
    actual = run(path, "rows", "3", "--out", output)
    assert {k: v for k, v in dry.items() if k != "dry_run"} == actual
    assert path.read_bytes() == source and output.exists()
    previous_out = output.read_bytes()
    seen = []
    original_check = insert.self_check

    def disk_check(staged, *args):
        assert isinstance(staged, Path) and staged.is_file()
        assert staged.parent == output.parent
        assert path.read_bytes() == source and output.read_bytes() == previous_out
        seen.append(True)
        return {"ok": False, "problems": ["injected failure"]}

    monkeypatch.setattr(insert, "self_check", disk_check)
    failed = run(path, "columns", "C", "--out", output)
    assert failed["status"] == "error" and seen
    assert path.read_bytes() == source and output.read_bytes() == previous_out
    assert not list(tmp_path.glob(".insert-*"))
    monkeypatch.setattr(insert, "self_check", original_check)


def test_ref_errors_external_and_pivot_warnings(tmp_path):
    path = basic(tmp_path / "errors.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws["H1"] = '=D4+D5+SUM(D4:D5)+#REF!+IF(TRUE,"#REF!",0)+[1]X!D4'
    wb["Other"]["A1"] = "=X!D4+'X'!D5+'[book.xlsx]X'!D4"
    wb.defined_names.add(DefinedName("Gone", attr_text="'X'!$D$4:$D$5"))
    wb.save(path)

    def patch(data):
        data["xl/pivotCache/pivotCacheDefinition1.xml"] = (
            f'<pivotCacheDefinition xmlns="{S}"><cacheSource type="worksheet"><worksheetSource ref="B3:D8" sheet="X"/></cacheSource></pivotCacheDefinition>'
        ).encode()

    patch_package(path, patch)
    before = payload(path)
    report = run(path, "delete-rows", "4", "--count", "2")
    assert report["status"] == "success", report
    after = payload(path)
    assert (
        after["xl/pivotCache/pivotCacheDefinition1.xml"]
        == before["xl/pivotCache/pivotCacheDefinition1.xml"]
    )
    errors = [w for w in report["warnings"] if w["type"] == "ref_error"]
    assert len(errors) == 7
    assert sum(w["sheet"] == "X" and w["cell"] == "H1" for w in errors) == 4
    assert sum(w["sheet"] == "Other" and w["cell"] == "A1" for w in errors) == 2
    assert (
        len([w for w in report["warnings"] if w["type"] == "external_reference"]) == 2
    )
    assert any(w["type"] == "pivot_cache_unchanged" for w in report["warnings"])
    assert insert.self_check(path, before, set(), "X", []) == {
        "ok": False,
        "problems": ["#REF! occurrence count does not match warnings"],
    }


def test_drop_complete_rules_and_filter_columns(tmp_path):
    path = basic(tmp_path / "drop.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws.merge_cells("D4:D5")
    ws.conditional_formatting.add("D3:D8", FormulaRule(formula=["D3>0"]))
    dv = DataValidation(type="list", formula1='"Yes,No"')
    dv.add("D3:D8")
    ws.add_data_validation(dv)
    ws["D4"].comment = Comment("Deleted", "Test")
    ws["D6"].hyperlink = "https://example.invalid"
    ws.auto_filter.ref = "B2:F9"
    ws.auto_filter.add_filter_column(2, ["1"])
    wb.save(path)
    report = run(path, "delete-columns", "D")
    assert report["status"] == "success", report
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    for name in ("mergeCells", "conditionalFormatting", "dataValidations"):
        assert root.find("s:" + name, NS) is None
    assert not root.findall("s:hyperlinks/s:hyperlink", NS)
    assert not root.findall("s:autoFilter/s:filterColumn", NS)
    assert len([w for w in report["warnings"] if w["type"] == "dropped_range"]) >= 4


def test_shared_master_deleted_and_nonuniform_translation(tmp_path):
    path = basic(tmp_path / "shared.xlsx")

    def patch(data):
        part = "xl/worksheets/sheet1.xml"
        root = xml(data, part)
        for row in range(3, 7):
            cell = cell_node(root, f"G{row}")
            cell[:] = []
            f = etree.SubElement(cell, f"{{{S}}}f", t="shared", si="7")
            if row == 3:
                f.set("ref", "G3:G6")
                f.text = "A3+$B$1"
        save_xml(data, part, root)

    patch_package(path, patch)
    report = run(path, "delete-rows", "3")
    assert report["status"] == "success", report
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    master = cell_node(root, "G3").find("s:f", NS)
    assert master.get("ref") == "G3:G5" and master.get("si") == "7"
    assert master.text == "A3+$B$1"
    out = load_workbook(path)["X"]
    assert out["G5"].value == "=A5+$B$1"

    # A follower referring two rows up crosses the insertion differently from
    # its master, so retaining one shared formula would silently change it.
    def nonuniform(data):
        part = "xl/worksheets/sheet1.xml"
        root = xml(data, part)
        cell_node(root, "G3").find("s:f", NS).text = "A1"
        save_xml(data, part, root)

    patch_package(path, nonuniform)
    report = run(path, "rows", "4")
    assert report["status"] == "success", report
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    assert not root.findall(".//s:f[@t='shared']", NS)
    assert formula(root, "G3") == "A1"
    assert formula(root, "G5") == "A2"
    assert formula(root, "G6") == "A3"
    assert any(w["type"] == "shared_formula_expanded" for w in report["warnings"])


def test_shared_formulas_on_other_sheet_need_individual_rewrites(tmp_path):
    path = basic(tmp_path / "shared_other.xlsx")
    wb = load_workbook(path)
    for row in range(1, 4):
        wb["Other"].cell(row, 1, "=X!A" + str(row + 2))
    wb.save(path)

    def patch(data):
        part = "xl/worksheets/sheet2.xml"
        root = xml(data, part)
        for row in range(1, 4):
            f = cell_node(root, f"A{row}").find("s:f", NS)
            f.set("t", "shared")
            f.set("si", "0")
            if row == 1:
                f.set("ref", "A1:A3")
            else:
                f.text = None
        save_xml(data, part, root)

    patch_package(path, patch)
    report = run(path, "delete-rows", "4")
    assert report["status"] == "success", report
    ws = load_workbook(path)["Other"]
    assert [ws[f"A{r}"].value for r in range(1, 4)] == ["=X!A3", "=X!#REF!", "=X!A4"]
    assert any(
        w["type"] == "ref_error" and w["sheet"] == "Other" and w["cell"] == "A2"
        for w in report["warnings"]
    )


def test_array_master_promotion(tmp_path):
    path = basic(tmp_path / "array.xlsx")

    def patch(data):
        part = "xl/worksheets/sheet1.xml"
        root = xml(data, part)
        cell = cell_node(root, "G3")
        cell[:] = []
        etree.SubElement(cell, f"{{{S}}}f", t="array", ref="G3:G6").text = "A3:A6*2"
        save_xml(data, part, root)

    patch_package(path, patch)
    report = run(path, "delete-rows", "3")
    assert report["status"] == "success", report
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    f = cell_node(root, "G3").find("s:f", NS)
    assert f.text == "A3:A5*2" and f.get("ref") == "G3:G5"
    assert any(w["type"] == "array_master_promoted" for w in report["warnings"])


@pytest.mark.parametrize(
    "op,at,count,expected",
    [
        ("columns", "B", 2, "D2:G8"),
        ("columns", "F", 2, "B2:E8"),
        ("columns", "D", 2, "B2:G8"),
        ("delete-columns", "B", 2, "B2:C8"),
        ("delete-columns", "D", 2, "B2:C8"),
        ("delete-rows", "2", 1, "B2:E7"),
    ],
)
def test_table_boundaries_and_multiple_columns(tmp_path, op, at, count, expected):
    path = basic(tmp_path / "table.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    for col, header in zip(range(2, 6), ["Column1", "column2", "Cost", "Profit"]):
        ws.cell(2, col, header)
    ws.add_table(Table(displayName="Sales", ref="B2:E8"))
    wb.save(path)
    report = run(path, op, at, "--count", str(count))
    assert report["status"] == "success", report
    table = xml(payload(path), "xl/tables/table1.xml")
    assert table.get("ref") == expected
    names = [c.get("name") for c in table.findall("s:tableColumns/s:tableColumn", NS)]
    if op == "columns" and at == "D":
        assert names == ["Column1", "column2", "Column3", "Column4", "Cost", "Profit"]
    if op == "delete-rows":
        assert names == ["302", "303", "304", "305"]
    ws = load_workbook(path)["X"]
    assert ws.tables["Sales"].ref == expected


def test_entire_table_removed_with_relationships(tmp_path):
    path = basic(tmp_path / "drop_table.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws["B2"], ws["C2"] = "One", "Two"
    ws.add_table(Table(displayName="Gone", ref="B2:C5"))
    wb.save(path)
    report = run(path, "delete-columns", "B", "--count", "2")
    assert report["status"] == "success", report
    after = payload(path)
    assert "xl/tables/table1.xml" not in after
    assert not xml(after, "xl/worksheets/sheet1.xml").findall("s:tableParts", NS)
    assert b'/table"' not in after["xl/worksheets/_rels/sheet1.xml.rels"]
    assert b"table1.xml" not in after["[Content_Types].xml"]
    assert not load_workbook(path)["X"].tables


@pytest.mark.parametrize(
    "args,message",
    [
        ([], "Expected"),
        (["sideways", "file.xlsx"], "Unknown operation"),
        (["rows", "file.xlsx", "--sheet", "X", "--at", "A"], "row numbers"),
        (["rows", "file.xlsx", "--sheet", "X", "--at", "0"], "row numbers"),
        (["columns", "file.xlsx", "--sheet", "X", "--at", "3"], "column letters"),
        (["columns", "file.xlsx", "--sheet", "X", "--at", "XFE"], "outside"),
        (["rows", "file.xlsx", "--sheet", "X", "--at", "1048577"], "outside"),
        (
            ["rows", "file.xlsx", "--sheet", "X", "--at", "3", "--count", "0"],
            "at least",
        ),
        (
            ["rows", "file.xlsx", "--sheet", "X", "--at", "3", "--count", "-1"],
            "at least",
        ),
        (
            ["rows", "file.xlsx", "--sheet", "X", "--at", "3", "--count", "1.5"],
            "integer",
        ),
        (["rows", "file.xlsx", "--sheet", "X", "--at"], "requires a value"),
        (["rows", "file.xlsx", "--sheet", "X"], "required"),
        (
            ["rows", "file.xlsx", "--sheet", "X", "--at", "3", "--wat"],
            "Unknown argument",
        ),
        (["rows", "file.xlsx", "--sheet", "X", "--at", "3", "--at", "4"], "Duplicate"),
        (
            [
                "delete-rows",
                "file.xlsx",
                "--sheet",
                "X",
                "--at",
                "3",
                "--copy-style-from",
                "2",
            ],
            "only valid",
        ),
    ],
)
def test_bad_arguments_one_json_error(args, message, capsys):
    assert insert.main(args) == 1
    report = json.loads(capsys.readouterr().out)
    assert report["status"] == "error" and message in report["message"]


def test_file_sheet_zip_and_used_range_errors(tmp_path, capsys):
    missing = tmp_path / "missing.xlsx"
    assert insert.main(["rows", str(missing), "--sheet", "X", "--at", "3"]) == 1
    assert json.loads(capsys.readouterr().out)["status"] == "error"
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not a zip")
    assert insert.main(["rows", str(invalid), "--sheet", "X", "--at", "3"]) == 1
    assert json.loads(capsys.readouterr().out)["status"] == "error"
    with zipfile.ZipFile(invalid, "w") as archive:
        archive.writestr("nothing.txt", "not a workbook")
    assert insert.main(["rows", str(invalid), "--sheet", "X", "--at", "3"]) == 1
    assert "Not a workbook" in json.loads(capsys.readouterr().out)["message"]
    path = basic(tmp_path / "input.xlsx")
    before = path.read_bytes()
    for args in [
        ["rows", str(path), "--sheet", "Missing", "--at", "3"],
        ["delete-rows", str(path), "--sheet", "X", "--at", "12", "--count", "2"],
        ["delete-columns", str(path), "--sheet", "X", "--at", "I"],
    ]:
        assert insert.main(args) == 1
        assert json.loads(capsys.readouterr().out)["status"] == "error"
        assert path.read_bytes() == before


@pytest.mark.parametrize(
    "damage,expected",
    [
        ("duplicate", "Duplicate cell"),
        ("rows", "Rows are unordered"),
        ("range", "Invalid"),
        ("formula", "Invalid formula"),
        ("name", "Invalid formula/name"),
        ("missing", "inventory"),
    ],
)
def test_self_check_detects_corruption(tmp_path, damage, expected):
    path = basic(tmp_path / "check.xlsx")
    original = payload(path)

    def corrupt(data):
        part = "xl/worksheets/sheet1.xml"
        root = xml(data, part)
        if damage == "duplicate":
            row = root.find("s:sheetData/s:row", NS)
            row.append(copy.deepcopy(row[0]))
        elif damage == "rows":
            rows = root.find("s:sheetData", NS)
            rows[:] = list(reversed(rows))
        elif damage == "range":
            etree.SubElement(root, f"{{{S}}}conditionalFormatting", sqref="A0:XFE4")
        elif damage == "formula":
            etree.SubElement(cell_node(root, "H1"), f"{{{S}}}f").text = '"unterminated'
        elif damage == "name":
            book = xml(data, "xl/workbook.xml")
            names = book.find("s:definedNames", NS)
            if names is None:
                names = etree.SubElement(book, f"{{{S}}}definedNames")
            etree.SubElement(
                names, f"{{{S}}}definedName", name="Bad"
            ).text = '"unterminated'
            save_xml(data, "xl/workbook.xml", book)
        elif damage == "missing":
            del data["docProps/core.xml"]
        save_xml(data, part, root)

    patch_package(path, corrupt)
    check = insert.self_check(path, original, set(), "X", [])
    assert not check["ok"]
    assert any(expected in problem for problem in check["problems"]), check


def test_help_and_cli_dry_run(tmp_path):
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    help_result = subprocess.run(
        [sys.executable, str(SCRIPT), "--help"],
        capture_output=True,
        text=True,
        env=env,
        cwd=ROOT,
    )
    assert (
        help_result.returncode == 0
        and help_result.stdout.strip() == insert.__doc__.strip()
    )
    assert not help_result.stderr
    path = basic(tmp_path / "cli.xlsx")
    before = path.read_bytes()
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "rows",
            str(path),
            "--sheet",
            "X",
            "--at",
            "3",
            "--dry-run",
        ],
        capture_output=True,
        text=True,
        env=env,
        cwd=ROOT,
    )
    assert result.returncode == 0 and not result.stderr
    report = json.loads(result.stdout)
    assert report["self_check"] == {"ok": True} and report["dry_run"] is True
    assert path.read_bytes() == before


def test_no_em_dashes():
    assert "\u2014" not in SCRIPT.read_text()


@pytest.mark.parametrize(
    "op,at,band,cf_formula,dv_formula,other_formula",
    [
        ("rows", "3", "D4:F9", "D4>0", "$D$4:$D$9", "X!D4>0"),
        ("columns", "C", "E3:G8", "E3>0", "$E$3:$E$8", "X!E3>0"),
        ("delete-rows", "4", "D3:F7", "D3>0", "$D$3:$D$7", "X!D3>0"),
        ("delete-columns", "D", "D3:E8", "#REF!>0", "#REF!", "X!#REF!>0"),
    ],
)
def test_rule_formulas_rewritten_in_target_and_other_sheets(
    tmp_path, op, at, band, cf_formula, dv_formula, other_formula
):
    path = basic(tmp_path / "rules.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws.conditional_formatting.add("D3:F8", FormulaRule(formula=["D3>0"]))
    dv = DataValidation(type="list", formula1="$D$3:$D$8")
    dv.add("D3:F8")
    ws.add_data_validation(dv)
    wb["Other"].conditional_formatting.add("A1:A8", FormulaRule(formula=["X!D3>0"]))
    wb.save(path)
    report = run(path, op, at)
    assert report["status"] == "success", report
    data = payload(path)
    root = xml(data, "xl/worksheets/sheet1.xml")
    cf = root.find("s:conditionalFormatting", NS)
    assert cf.get("sqref") == band
    assert cf.findtext("s:cfRule/s:formula", namespaces=NS) == cf_formula
    dv = root.find("s:dataValidations/s:dataValidation", NS)
    assert (
        dv.get("sqref") == band
        and dv.findtext("s:formula1", namespaces=NS) == dv_formula
    )
    cf = xml(data, "xl/worksheets/sheet2.xml").find("s:conditionalFormatting", NS)
    assert cf.get("sqref") == "A1:A8"
    assert cf.findtext("s:cfRule/s:formula", namespaces=NS) == other_formula
    if op == "delete-columns":
        errors = [w for w in report["warnings"] if w["type"] == "ref_error"]
        assert len(errors) == 3
        assert {w["sheet"] for w in errors} == {"X", "Other"}


def test_extension_rules_and_sparklines_drop_only_deleted_items(tmp_path):
    path = rich(tmp_path / "ext.xlsx")

    def patch(data):
        part = "xl/worksheets/sheet2.xml"
        root = xml(data, part)
        extlist = root.find("s:extLst", NS)
        ext = etree.SubElement(extlist, f"{{{S}}}ext", uri="{dv-test}")
        validations = etree.SubElement(ext, f"{{{X14}}}dataValidations", count="2")
        for band in ("D3:D8", "G3:G8"):
            dv = etree.SubElement(validations, f"{{{X14}}}dataValidation", type="list")
            etree.SubElement(
                etree.SubElement(dv, f"{{{X14}}}formula1"), f"{{{XM}}}f"
            ).text = "X!B3:B8"
            etree.SubElement(dv, f"{{{XM}}}sqref").text = band
        root.find(".//x14:sparkline/xm:sqref", NS).text = "D5"
        save_xml(data, part, root)

    patch_package(path, patch)
    report = run(path, "delete-columns", "D")
    assert report["status"] == "success", report
    root = xml(payload(path))
    validations = root.find(".//x14:dataValidations", NS)
    assert validations.get("count") == "1" and len(validations) == 1
    assert validations[0].findtext("xm:sqref", namespaces=NS) == "F3:F8"
    assert not root.findall(".//x14:sparklineGroup", NS)
    assert any(w["type"] == "sparkline_dropped" for w in report["warnings"])


def test_freeze_removal_and_pure_split_units(tmp_path):
    path = basic(tmp_path / "views.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws.freeze_panes = "D4"
    ws.sheet_view.selection[-1].activeCell = "D4"
    ws.sheet_view.selection[-1].sqref = "D4"
    wb.save(path)
    report = run(path, "delete-columns", "A", "--count", "3")
    assert report["status"] == "success", report
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    pane = root.find("s:sheetViews/s:sheetView/s:pane", NS)
    assert pane.get("xSplit") is None and pane.get("ySplit") == "3"
    assert pane.get("topLeftCell") == "A4" and pane.get("activePane") == "bottomLeft"
    selections = root.findall("s:sheetViews/s:sheetView/s:selection", NS)
    assert [s.get("pane") for s in selections] == [None, "bottomLeft"]
    report = run(path, "delete-rows", "1", "--count", "3")
    assert report["status"] == "success", report
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    assert root.find("s:sheetViews/s:sheetView/s:pane", NS) is None
    assert len(root.findall("s:sheetViews/s:sheetView/s:selection", NS)) == 1

    def split(data):
        part = "xl/worksheets/sheet1.xml"
        root = xml(data, part)
        view = root.find("s:sheetViews/s:sheetView", NS)
        view.insert(
            0,
            etree.Element(
                f"{{{S}}}pane",
                state="split",
                xSplit="2400.5",
                ySplit="1200.5",
                topLeftCell="D4",
            ),
        )
        save_xml(data, part, root)

    patch_package(path, split)
    assert run(path, "rows", "3")["status"] == "success"
    pane = xml(payload(path), "xl/worksheets/sheet1.xml").find(
        "s:sheetViews/s:sheetView/s:pane", NS
    )
    assert pane.get("xSplit") == "2400.5" and pane.get("ySplit") == "1200.5"
    assert pane.get("topLeftCell") == "D5"


@pytest.mark.parametrize("op,at", [("rows", "30"), ("columns", "Z")])
def test_unaffected_attached_parts_byte_identical(tmp_path, op, at):
    path = rich(tmp_path / "far.xlsx")

    def compact_anchor(data):
        part = "xl/drawings/commentsDrawing1.vml"
        root = xml(data, part)
        root.find(".//x:Anchor", NS).text = "3,15,4,2,6,20,8,4"
        save_xml(data, part, root)

    patch_package(path, compact_anchor)
    before = payload(path)
    report = run(path, op, at)
    assert report["status"] == "success", report
    after = payload(path)
    assert_preserved(before, after, report)
    for part in (
        "xl/drawings/drawing1.xml",
        "xl/drawings/commentsDrawing1.vml",
        "xl/charts/chart1.xml",
        "xl/comments/comment1.xml",
        "xl/threadedComments/threadedComment1.xml",
        "xl/tables/table1.xml",
        "xl/worksheets/sheet2.xml",
    ):
        assert before[part] == after[part], part


def test_table_calculated_formulas_and_internal_hyperlinks(tmp_path):
    path = basic(tmp_path / "additional.xlsx")
    wb = load_workbook(path)
    ws = wb["X"]
    ws["B2"], ws["C2"] = "One", "Two"
    ws.add_table(Table(displayName="Sales", ref="B2:C8"))
    wb.save(path)

    def patch(data):
        part = "xl/tables/table1.xml"
        table = xml(data, part)
        columns = table.findall("s:tableColumns/s:tableColumn", NS)
        etree.SubElement(
            columns[0], f"{{{S}}}calculatedColumnFormula"
        ).text = "D4+Sales[Two]"
        etree.SubElement(columns[1], f"{{{S}}}totalsRowFormula").text = "SUM(X!D3:D8)"
        save_xml(data, part, table)
        part = "xl/worksheets/sheet2.xml"
        root = xml(data, part)
        links = etree.SubElement(root, f"{{{S}}}hyperlinks")
        etree.SubElement(links, f"{{{S}}}hyperlink", ref="A1", location="X!D4")
        etree.SubElement(links, f"{{{S}}}hyperlink", ref="A2", location="#X!D8")
        save_xml(data, part, root)

    patch_package(path, patch)
    report = run(path, "delete-rows", "4")
    assert report["status"] == "success", report
    data = payload(path)
    table = xml(data, "xl/tables/table1.xml")
    assert (
        table.findtext(".//s:calculatedColumnFormula", namespaces=NS)
        == "#REF!+Sales[Two]"
    )
    assert table.findtext(".//s:totalsRowFormula", namespaces=NS) == "SUM(X!D3:D7)"
    links = xml(data, "xl/worksheets/sheet2.xml").findall(
        "s:hyperlinks/s:hyperlink", NS
    )
    assert [link.get("location") for link in links] == ["X!#REF!", "#X!D7"]
    assert len([w for w in report["warnings"] if w["type"] == "ref_error"]) == 2


@pytest.mark.parametrize("op,at", [("rows", "1048576"), ("columns", "XFD")])
def test_grid_overflow_fails_without_writing(tmp_path, op, at):
    path = tmp_path / "edge.xlsx"
    wb = Workbook()
    wb.active.title = "X"
    wb.active["XFD1048576"] = 1
    wb.save(path)
    before = path.read_bytes()
    with pytest.raises(ValueError, match="outside Excel"):
        run(path, op, at)
    assert path.read_bytes() == before


def test_empty_sheet_and_entire_used_band_deletion(tmp_path):
    path = tmp_path / "empty.xlsx"
    wb = Workbook()
    wb.active.title = "X"
    wb.save(path)
    assert run(path, "rows", "3")["self_check"]["ok"]
    assert load_workbook(path)["X"].max_row == 1
    path = basic(tmp_path / "all.xlsx", rows=4, columns=3)
    assert run(path, "delete-rows", "1", "--count", "4")["self_check"]["ok"]
    root = xml(payload(path), "xl/worksheets/sheet1.xml")
    assert not root.findall("s:sheetData/s:row", NS)
    assert root.find("s:dimension", NS).get("ref") == "A1"


def test_defined_name_expression_validation(tmp_path):
    path = basic(tmp_path / "name.xlsx")
    original = payload(path)
    wb = load_workbook(path)
    wb.defined_names.add(DefinedName("BadRange", attr_text="'X'!$B$0:$D$5"))
    wb.defined_names.add(DefinedName("BadFunction", attr_text="SUM(X!A1"))
    wb.save(path)
    check = insert.self_check(path, original, set(), "X", [])
    assert not check["ok"]
    assert len([p for p in check["problems"] if "Invalid formula/name" in p]) == 2


def test_saved_active_sheet_for_workbook_names(tmp_path):
    path = basic(tmp_path / "names.xlsx")
    wb = load_workbook(path)
    wb.defined_names.add(
        DefinedName("Dynamic", attr_text="OFFSET($B$3,0,0,COUNTA(B:B),1)")
    )
    wb.save(path)
    report = run(path, "rows", "3")
    assert report["status"] == "success", report
    assert (
        load_workbook(path).defined_names["Dynamic"].attr_text
        == "OFFSET($B$4,0,0,COUNTA(B:B),1)"
    )


def test_unreadable_input_and_failed_replace_leave_files_intact(
    tmp_path, monkeypatch, capsys
):
    path = basic(tmp_path / "unreadable.xlsx")
    before = path.read_bytes()
    real_zip = insert.zipfile.ZipFile

    def unreadable(*args, **kwargs):
        raise PermissionError("simulated unreadable workbook")

    monkeypatch.setattr(insert.zipfile, "ZipFile", unreadable)
    assert insert.main(["rows", str(path), "--sheet", "X", "--at", "3"]) == 1
    assert "PermissionError" in json.loads(capsys.readouterr().out)["message"]
    monkeypatch.setattr(insert.zipfile, "ZipFile", real_zip)

    def fail_replace(*args):
        raise OSError("simulated replace failure")

    monkeypatch.setattr(insert.os, "replace", fail_replace)
    assert insert.main(["rows", str(path), "--sheet", "X", "--at", "3"]) == 1
    assert "replace failure" in json.loads(capsys.readouterr().out)["message"]
    assert path.read_bytes() == before
    assert not list(tmp_path.glob(".insert-*"))


@pytest.mark.parametrize(
    "op,shifted,rewritten", [("rows", 4, 3), ("delete-rows", 2, 2)]
)
def test_exact_report_counts_and_success_exit_for_ref_errors(
    tmp_path, capsys, op, shifted, rewritten
):
    path = tmp_path / "counts.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "X"
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, 3
    ws["B2"] = "=A2"
    ws["B3"] = "=SUM(A1:A3)"
    wb.create_sheet("Other")["A1"] = "=X!B3"
    wb.defined_names.add(DefinedName("Data", attr_text="X!A2"))
    wb.save(path)

    def no_calc(data):
        part = "xl/workbook.xml"
        root = xml(data, part)
        root.remove(root.find("s:calcPr", NS))
        save_xml(data, part, root)

    patch_package(path, no_calc)
    assert insert.main([op, str(path), "--sheet", "X", "--at", "2"]) == 0
    report = json.loads(capsys.readouterr().out)
    assert report["cells_shifted"] == shifted
    assert report["formulas_rewritten"] == rewritten
    assert report["names_rewritten"] == 1
    assert report["charts_rewritten"] == 0
    assert report["self_check"] == {"ok": True}
    if op == "delete-rows":
        assert len([w for w in report["warnings"] if w["type"] == "ref_error"]) == 1
    assert (
        xml(payload(path), "xl/workbook.xml").find("s:calcPr", NS).get("fullCalcOnLoad")
        == "1"
    )
